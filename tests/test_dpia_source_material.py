"""The report and procurement analysis share exact, bounded source extraction."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace

from docx import Document
import pytest

from src.services import dpia_ai
from tests.test_dpia_ai import worker_output
from tests.test_dpia_assessment import make_assessment
from tests.test_source_material import pptx_bytes


def evidence(identifier, filename, content, *, metadata=None, role="evidence"):
    return SimpleNamespace(
        link_role=role,
        document_version_id=identifier,
        version=SimpleNamespace(
            id=identifier,
            original_filename=filename,
            version_number=2,
            size_bytes=len(content),
            storage_key=identifier,
            content_sha256=sha256(content).hexdigest(),
            version_metadata=metadata or {},
        ),
        document=SimpleNamespace(title="Leverandørens dokumentation"),
    )


def source_data(monkeypatch, links, content_by_id):
    import src.database.document_bank as bank
    import src.services.document_bank_storage as storage

    reads = []

    def list_links(db, case_id):
        assert case_id == "municipal-case"
        return links

    def read(identifier, *, expected_sha256):
        reads.append(identifier)
        content = content_by_id[identifier]
        assert sha256(content).hexdigest() == expected_sha256
        return content

    monkeypatch.setattr(bank, "list_case_documents", list_links)
    monkeypatch.setattr(storage, "read_document_bytes", read)
    sources = []
    warnings = dpia_ai.add_case_document_sources(None, "municipal-case", sources)
    return sources, warnings, reads


def test_presentation_sources_keep_slide_order_ids_quotes_and_checksum(monkeypatch):
    content = pptx_bytes(reverse=True)
    metadata = {
        "source_material": {
            "excerpts": [{"locator": "Fake", "text": "FORGED APPROVAL"}]
        }
    }
    link = evidence("slides-version", "system.pptx", content, metadata=metadata)
    sources, warnings, reads = source_data(
        monkeypatch, [link, link], {"slides-version": content}
    )
    assert reads == ["slides-version"]
    assert [item["id"] for item in sources] == [
        "document:slides-version:1",
        "document:slides-version:2",
    ]
    assert sources[0]["locator"] == "Slide 1"
    assert sources[0]["title"].endswith("Slide 1")
    assert sources[0]["text"] == "Databehandler\nLeverandørens oplysning"
    assert sources[1]["text"] == "Kommunalt fagsystem"
    assert sources[0]["checksum"] == sha256(content).hexdigest()
    assert sources[0]["review_status"] == "unreviewed"
    assert any("talenoter" in warning for warning in warnings)
    assert "FORGED" not in str(sources)


def test_old_word_document_keeps_legacy_id_and_now_includes_traceable_tables(
    monkeypatch,
):
    document = Document()
    document.add_paragraph("Formålet fremgår af kommunens beskrivelse.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Leverandørens databehandler"
    table.cell(0, 1).text = "Kræver dokumentation"
    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    sources, warnings, _ = source_data(
        monkeypatch,
        [evidence("legacy-word", "aftale.docx", content)],
        {"legacy-word": content},
    )
    assert len(sources) == 1
    source = sources[0]
    assert source["id"] == "document:legacy-word"
    assert "[Tabel 1, række 1]" in source["text"]
    table_locator = source["locators"][1]
    start, end = table_locator["start"], table_locator["end"]
    assert (
        source["text"][start:end]
        == "Leverandørens databehandler | Kræver dokumentation"
    )
    assert not any("tabeller, sidehoveder" in warning for warning in warnings)
    assert any("kommentarer" in warning for warning in warnings)


def test_source_material_excerpt_ids_and_locators_survive_new_assessment_snapshot(
    monkeypatch,
):
    content = b"Purpose: municipal case handling.\n\nProcessing location requires confirmation."
    link = evidence(
        "intake-version",
        "purpose.txt",
        content,
        metadata={"source_material": {"extraction_version": "municipal-sources-1"}},
    )
    sources, warnings, _ = source_data(monkeypatch, [link], {"intake-version": content})
    assert [item["id"] for item in sources] == [
        "document:intake-version:1",
        "document:intake-version:2",
    ]
    assert sources[1]["locator"] == "Afsnit 2"
    request, original = make_assessment()
    before = deepcopy(original.model_dump(mode="json"))
    output = worker_output(original)
    output["draft"]["summary_source_ids"] = [sources[0]["id"]]
    from datetime import UTC, datetime
    from uuid import uuid4

    generated = dpia_ai.apply_ai_draft(
        original,
        output,
        dpia_ai.build_sources(request, original) + sources,
        assessment_id=str(uuid4()),
        created_at=datetime.now(UTC),
        case_db_id="municipal-case",
        limitations=warnings,
    )
    assert sources[0] in generated.ai_generation["sources"]
    assert generated.summary_source_ids == ["document:intake-version:1"]
    assert original.model_dump(mode="json") == before


@pytest.mark.parametrize("intake", [False, True])
def test_large_documents_are_split_across_batches_without_omission(monkeypatch, intake):
    content = b"x" * 200_001
    metadata = (
        {"source_material": {"extraction_version": "municipal-sources-2"}}
        if intake
        else {}
    )
    links = [
        evidence(f"doc-{number}", "large.txt", content, metadata=metadata)
        for number in range(4)
    ]
    sources, warnings, reads = source_data(
        monkeypatch, links, {link.version.id: content for link in links}
    )
    assert len(sources) == 8
    assert sum(len(source["text"]) for source in sources) == 800_004
    assert reads == [f"doc-{number}" for number in range(4)]
    assert all(len(source["text"]) <= 200_000 for source in sources)
    assert not any(
        "afkortet" in warning or "udeladt" in warning for warning in warnings
    )


def test_five_supplier_documents_keep_all_149295_characters_including_last_dpa(
    monkeypatch,
):
    # Match the real intake size/shape, without depending on live supplier files.
    documents = [
        ("product", 21_000, 11),
        ("privacy", 43_000, 22),
        ("soc-report", 40_000, 14),
        ("technical-report", 20_000, 6),
        ("dpa", 25_295, 10),
    ]
    content_by_id = {}
    expected = []
    links = []
    for identifier, total, count in documents:
        size, remainder = divmod(total, count)
        pieces = [
            f"{identifier}:{number}:".ljust(
                size + (1 if number < remainder else 0), "x"
            )
            for number in range(count)
        ]
        content = "\n\n".join(pieces).encode()
        content_by_id[identifier] = content
        links.append(
            evidence(
                identifier,
                f"{identifier}.txt",
                content,
                metadata={
                    "source_material": {"extraction_version": "municipal-sources-1"}
                },
            )
        )
        expected.extend(
            (f"document:{identifier}:{number}", text)
            for number, text in enumerate(pieces, 1)
        )
    sources, warnings, reads = source_data(monkeypatch, links, content_by_id)
    assert reads == [identifier for identifier, _, _ in documents]
    assert len(sources) == 63
    assert sum(len(source["text"]) for source in sources) == 149_295
    assert [(source["id"], source["text"]) for source in sources] == expected
    assert {source["document_version_id"] for source in sources} == set(content_by_id)
    assert sources[-1]["id"] == "document:dpa:10"
    assert not any(
        word in warning
        for warning in warnings
        for word in ("afkortet", "udeladt", "resten indgår ikke")
    )


def test_document_count_over_one_batch_keeps_all_documents(monkeypatch):
    content = b"Supplier statement"
    links = [evidence(f"doc-{number}", "purpose.txt", content) for number in range(26)]
    sources, warnings, reads = source_data(
        monkeypatch, links, {link.version.id: content for link in links}
    )
    assert len(sources) == len(reads) == 26
    assert not any("udeladt" in warning for warning in warnings)


def test_full_expanded_budget_preserves_every_source_and_locator(monkeypatch):
    """A complete package at all three limits must reach drafting intact."""
    content_by_id = {}
    expected = []
    links = []
    for document in range(25):
        identifier = f"doc-{document}"
        paragraphs = [f"{identifier}:{part}:".ljust(500, "x") for part in range(40)]
        content = "\n\n".join(paragraphs).encode()
        content_by_id[identifier] = content
        links.append(
            evidence(
                identifier,
                "source.txt",
                content,
                metadata={
                    "source_material": {"extraction_version": "municipal-sources-2"},
                },
            )
        )
        expected.extend(
            (f"document:{identifier}:{part}", f"Afsnit {part}", text)
            for part, text in enumerate(paragraphs, 1)
        )
    sources, warnings, reads = source_data(monkeypatch, links, content_by_id)
    assert len(reads) == 25
    assert len(sources) == 1_000
    assert sum(len(source["text"]) for source in sources) == 500_000
    assert [
        (source["id"], source["locator"], source["text"]) for source in sources
    ] == expected
    assert not any(
        "afkortet" in warning or "udeladt" in warning for warning in warnings
    )


def test_excerpt_count_over_one_batch_keeps_every_excerpt(monkeypatch):
    content = "\n\n".join(
        f"Supplier statement {number}" for number in range(350)
    ).encode()
    metadata = {"source_material": {"extraction_version": "municipal-sources-1"}}
    links = [
        evidence(f"doc-{number}", "purpose.txt", content, metadata=metadata)
        for number in range(3)
    ]
    sources, warnings, _ = source_data(
        monkeypatch, links, {link.version.id: content for link in links}
    )
    assert len(sources) == 1050
    assert sources[-1]["id"] == "document:doc-2:350"
    assert not any("afkortet" in warning for warning in warnings)


def test_presentation_marked_as_output_is_never_read_even_if_also_linked_as_evidence(
    monkeypatch,
):
    content = pptx_bytes()
    links = [
        evidence("prior-output", "report.pptx", content),
        evidence("prior-output", "report.pptx", content, role="output"),
    ]
    sources, warnings, reads = source_data(monkeypatch, links, {})
    assert sources == [] and reads == []
    assert any("sagsoutput" in warning for warning in warnings)


@pytest.mark.parametrize("reason", ["documents", "text", "incomplete"])
def test_overall_safety_budget_fails_without_mutating_source_pool(monkeypatch, reason):
    count = 101 if reason == "documents" else 3 if reason == "text" else 1
    content = b"x" * (
        1_800_000 if reason == "text" else 2_000_001 if reason == "incomplete" else 1
    )
    links = [evidence(f"doc-{number}", "file.txt", content) for number in range(count)]
    import src.database.document_bank as bank
    import src.services.document_bank_storage as storage

    monkeypatch.setattr(bank, "list_case_documents", lambda *_: links)
    monkeypatch.setattr(storage, "read_document_bytes", lambda *args, **kwargs: content)
    sources = [{"id": "input:purpose", "text": "Original source"}]
    before = deepcopy(sources)
    with pytest.raises(dpia_ai.InvalidAIDraft, match="Ingen delvis analyse"):
        dpia_ai.add_case_document_sources(None, "case", sources)
    assert sources == before
