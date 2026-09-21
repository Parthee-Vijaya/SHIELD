"""Reading aids may clarify a decision, but cannot silently change it."""

from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
import pytest

from src.services.dpia_assessment import DPIARecommendation
from src.services.dpia_docx import export_dpia_docx
from src.services.dpia_export import TEMPLATE_PATH, export_dpia_xlsx
from src.services.dpia_reading_guide import reading_guide
from tests.test_dpia_assessment import make_assessment
from tests.test_dpia_docx import text_content


@pytest.mark.parametrize("provider", ["codex-local", "codex-local-test"])
def test_exports_show_actual_codex_origin_and_preserve_test_label_only_for_test_provider(
    provider,
):
    request, result = make_assessment()
    result.ai_generation = {
        "provider": provider,
        "model": "gpt-6-astra",
        "run_id": "municipal-scenario-2026-09-21",
        "model_run_provenance": "operator_reported",
        "review": {"model": "typesafe-ai/jev", "checks": []},
    }
    before = deepcopy(result.model_dump(mode="json"))
    texts = [text_content(Document(BytesIO(export_dpia_docx(request, result))))]
    with ZipFile(BytesIO(export_dpia_xlsx(request, result))) as archive:
        texts.append(
            " ".join(
                ET.fromstring(
                    archive.read("xl/worksheets/ai-provenance.xml")
                ).itertext()
            )
        )
    for text in texts:
        assert "gpt-6-astra" in text
        assert "typesafe-ai/jev" in text
        assert "municipal-scenario-2026-09-21" in text
        assert (
            "Model og kørsels-ID er oplyst af den lokale operatør ved importen." in text
        )
        assert "openai/gpt-5.5" not in text
        assert "Udarbejdet i AI Gateway" not in text
        if provider == "codex-local":
            assert "Codex – lokal kørsel" in text
            assert "Kørsels-ID" in text
            assert "Codex – lokal testkørsel" not in text
            assert "Testkørsel" not in text
        else:
            assert "Codex – lokal testkørsel" in text
            assert "Testkørsel" in text
    assert result.model_dump(mode="json") == before


def test_exports_separate_saved_recommendations_without_changing_assessment():
    request, result = make_assessment()
    result.recommendations = [
        DPIARecommendation(
            id="local-model",
            title="Afprøv lokal bearbejdning",
            proposal="Undersøg en lokal model til indledende bearbejdning.",
            rationale="Det kan begrænse videregivelsen af input.",
            prerequisites="Afklar licens, drift og dansk kvalitet før et valg.",
            verification="Test kvalitet og netværkstrafik med syntetiske data.",
            source_ids=["input:purpose"],
        )
    ]
    result.executive_summary = (
        "Første korte afsnit.\n\n- Afklar hjemlen.\n- Afprøv sletning."
    )
    before = deepcopy(result.model_dump(mode="json"))
    word = Document(BytesIO(export_dpia_docx(request, result)))
    text = text_content(word)
    assert "Beslutningsoversigt" in text
    assert "Anbefalinger – forslag til faglig drøftelse" in text
    assert "Gemte forslag i denne vurderingsversion." in text
    assert (
        "De dokumenterer ikke gennemførte foranstaltninger, godkendelse eller lavere risiko."
        in text
    )
    assert "Afklar licens, drift og dansk kvalitet før et valg." in text
    assert "Hvad kan ske?" in text
    assert "Hvorfor er dette en risiko?" in text
    assert "Hvem rammes – og hvad er konsekvensen?" in text
    assert "Mitigerende forslag og oplyste foranstaltninger" in text
    assert any(
        p.text == "Afklar hjemlen." and p.style.name == "List Bullet"
        for p in word.paragraphs
    )
    with ZipFile(BytesIO(export_dpia_xlsx(request, result))) as output, ZipFile(
        TEMPLATE_PATH
    ) as original:
        supplement = ET.fromstring(output.read("xl/worksheets/ai-provenance.xml"))
        supplement_text = " ".join(supplement.itertext())
        assert "Anbefalinger – forslag til faglig drøftelse" in supplement_text
        assert "Gemte forslag i denne vurderingsversion." in supplement_text
        assert "Test kvalitet og netværkstrafik med syntetiske data." in supplement_text
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        risks = ET.fromstring(output.read("xl/worksheets/sheet6.xml"))
        assert "Hvorfor er dette en risiko?" in " ".join(risks.itertext())
        assert [f.text for f in risks.findall(".//m:f", ns)] == [
            f.text
            for f in ET.fromstring(original.read("xl/worksheets/sheet6.xml")).findall(
                ".//m:f", ns
            )
        ]
        assert not supplement.findall(".//m:f", ns)
    assert result.model_dump(mode="json") == before


def test_version_bound_decision_is_consistent_in_word_and_official_excel_section():
    request, result = make_assessment()
    result.reading_guide = reading_guide(
        request,
        result,
        approvals=[
            {
                "id": "decision-1",
                "status": "approved_with_conditions",
                "approval_type": "dpia",
                "decided_by": "Syntetisk jurist",
                "decided_at": "2026-09-21T13:00:00Z",
                "is_identity_verified": True,
                "subject_reference_type": "dpia",
                "subject_reference_id": result.id,
                "reason": "Den afgrænsede pilot er gennemgået.",
                "conditions": "Kun syntetiske data må anvendes.",
            }
        ],
    )
    word = text_content(Document(BytesIO(export_dpia_docx(request, result))))
    with ZipFile(BytesIO(export_dpia_xlsx(request, result))) as archive:
        stakeholders = " ".join(
            ET.fromstring(archive.read("xl/worksheets/sheet8.xml")).itertext()
        )
    for text in (word, stakeholders):
        assert "Godkendt med vilkår" in text
        assert "Syntetisk jurist" in text
        assert "Kun syntetiske data må anvendes." in text
        assert "Beslutningstype: Konsekvensanalyse" in text
        assert (
            "Formel ledelsesgodkendelse og accept af resterende risiko er ikke dokumenteret"
            not in text
        )


def test_historical_exports_label_new_reading_aids_and_do_not_infer_approval():
    request, result = make_assessment(solution_type="saas")
    result.recommendations = []
    result.status = "ready_for_review"
    result.status_label = "Klar til faglig gennemgang"
    result.blockers = ["Den konkrete behandlingshjemmel skal dokumenteres."]
    result.ai_generation = {"review": {"status": "passed", "checks": []}}
    before = deepcopy(result.model_dump(mode="json"))
    outputs = [text_content(Document(BytesIO(export_dpia_docx(request, result))))]
    with ZipFile(BytesIO(export_dpia_xlsx(request, result))) as archive:
        outputs.append(
            " ".join(
                ET.fromstring(
                    archive.read("xl/worksheets/ai-provenance.xml")
                ).itertext()
            )
        )
    for text in outputs:
        assert "Regelbaserede forslag ud fra sagens oplysninger" in text
        assert "indgår ikke i det historiske AI-udkast eller dets JEV-kontrol" in text
        assert "Godkendelse fremgår ikke af denne rapportversion" in text
        assert "lokal behandling" in text
        assert "sletning" in text
        assert "Godkendelse registreret" not in text
        assert result.blockers[0] in text
    assert result.model_dump(mode="json") == before


def test_unknown_scale_transfer_dpo_and_human_control_remain_explicit_in_exports():
    request, result = make_assessment(
        large_scale=None,
        transfer_outside_eea=None,
        transfer_mechanism="not_assessed",
        dpo_involved=None,
        human_oversight=None,
    )
    word = text_content(Document(BytesIO(export_dpia_docx(request, result))))
    assert "Ikke afklaret. Kriteriet medregnes forsigtigt" in word
    with ZipFile(BytesIO(export_dpia_xlsx(request, result))) as archive:
        excel = " ".join(
            " ".join(ET.fromstring(archive.read(name)).itertext())
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    for content in (word, excel):
        assert "inddragelse er ikke dokumenteret" in content
        assert "Det er ikke afklaret, om personoplysninger overføres" in content
        assert "Foreløbig screening" in content
        assert "Menneskelig kontrol af AI-output er ikke afklaret" in content
        assert "DPO er ikke oplyst som inddraget" not in content
