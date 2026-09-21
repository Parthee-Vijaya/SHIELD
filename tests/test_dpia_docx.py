"""Word exports must preserve the assessment, provenance and review boundary."""

from datetime import UTC, datetime
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
import pytest

from src.services.dpia_assessment import DPIAAssessmentRequest, assess_dpia
from src.services.dpia_docx import export_dpia_docx
from src.services.dpia_export import RISK_ROWS, SECTION_ROWS, TemplateError


@pytest.fixture
def assessment():
    request = DPIAAssessmentRequest(
        project_name="Syntetisk dokumentassistent",
        organisation="Eksempelkommune",
        owner="Systemejer for dokumentassistent",
        purpose="At hjælpe medarbejdere med at finde relevante interne vejledninger.",
        processing_description=(
            "Medarbejdere søger i interne vejledninger og får et kildebaseret "
            "tekstforslag. En medarbejder gennemgår altid svaret før brug."
        ),
        data_subjects=["employees"],
        personal_data_categories=["identity", "usage_data"],
        special_categories=False,
        criminal_data=False,
        vulnerable_subjects=False,
        large_scale=False,
        systematic_monitoring=False,
        automated_decisions=False,
        human_oversight=True,
        solution_type="ai_system",
        supplier_name="Syntetisk Leverandør",
        hosting_region="eu_eea",
        transfer_outside_eea=False,
        transfer_mechanism="not_applicable",
        model_training=False,
        legal_basis="not_assessed",
        dpo_involved=True,
        controls=["access_control", "testing"],
        verified_controls=["access_control"],
        control_evidence={
            "access_control": "Syntetisk adgangstest AC-01 er gennemført."
        },
    )
    result = assess_dpia(
        request,
        assessment_id="synthetic-docx-v1",
        created_at=datetime(2026, 9, 20, 12, 0, tzinfo=UTC),
    )
    return request, result


def text_content(doc):
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ]
    )


def test_word_contains_every_section_risk_and_review_issue(assessment):
    request, result = assessment
    doc = Document(BytesIO(export_dpia_docx(request, result)))
    text = text_content(doc)
    headings = {p.text for p in doc.paragraphs if p.style.name.startswith("Heading")}
    assert len(result.sections) == 39 and len(result.risks) == 33
    assert all(
        f"{s.id} {s.title}" in headings and s.text in text for s in result.sections
    )
    assert all(
        f"{r.id} {r.area}" in headings
        and r.scenario in text
        and r.measures in text
        and r.owner in text
        for r in result.risks
    )
    assert all(
        value in text
        for value in result.blockers + result.missing_information + result.next_steps
    )
    assert result.scope in text and result.executive_summary in text
    assert "Velatir" not in text and "Scaleway" not in text
    assert {s.id for s in result.sections} == set(SECTION_ROWS)
    assert {r.id for r in result.risks} == set(RISK_ROWS)


def test_word_uses_exact_matrix_and_before_after_scores(assessment):
    request, result = assessment
    doc = Document(BytesIO(export_dpia_docx(request, result)))
    matrix = next(
        t for t in doc.tables if t.cell(0, 0).text == "Konsekvens / sandsynlighed"
    )
    assert [[c.text for c in row.cells[1:]] for row in matrix.rows[1:]] == [
        ["Høj", "Meget høj", "Meget høj", "Meget høj"],
        ["Mellem", "Høj", "Høj", "Meget høj"],
        ["Lav", "Mellem", "Høj", "Høj"],
        ["Lav", "Lav", "Mellem", "Høj"],
    ]
    tables = [t for t in doc.tables if t.cell(0, 0).text == "Vurdering"]
    assert len(tables) == 33
    for table, risk in zip(tables, result.risks):
        assert f"({risk.likelihood})" in table.cell(1, 1).text
        assert f"({risk.impact})" in table.cell(1, 2).text
        assert f"({risk.residual_likelihood})" in table.cell(2, 1).text
        assert f"({risk.residual_impact})" in table.cell(2, 2).text


def test_word_shows_actual_model_and_limits_without_platform_tag(assessment):
    request, result = assessment
    result.ai_generation = {
        "provider": "codex-local-test",
        "model": "gpt-5.6-sol",
        "run_id": "synthetic-codex-run-001",
        "review": {"model": "typesafe-ai/jev", "checks": []},
        "limitations": ["Ingen dokumentbankfiler indgår i denne test."],
    }
    text = text_content(Document(BytesIO(export_dpia_docx(request, result))))
    assert "GPT-5.6 Sol" in text
    assert "typesafe-ai/jev" in text
    assert "codex" not in text.lower()
    assert "Udarbejdet i" not in text
    assert "Ingen dokumentbankfiler indgår i denne test." in text
    assert "openai/gpt-5.5" not in text


def test_word_does_not_turn_dpo_involvement_or_ai_review_into_approval(assessment):
    request, result = assessment
    result = result.model_copy(
        update={
            "status": "ready_for_review",
            "status_label": "Klar til faglig gennemgang",
            "ai_generation": {
                "model": "example/model",
                "review": {"status": "passed"},
                "approval": {"approved": True, "approver": "AI-FAKE-APPROVER"},
            },
        }
    )
    text = text_content(Document(BytesIO(export_dpia_docx(request, result))))
    assert "konkrete synspunkter er ikke dokumenteret" in text
    assert (
        "Formel ledelsesgodkendelse og accept af resterende risiko er ikke dokumenteret"
        in text
    )
    assert "Vurderingsstatus er ikke en godkendelse" in text
    assert "AI-FAKE-APPROVER" not in text
    assert "passed" in text


def test_word_keeps_version_sources_additional_proposals_and_questions(assessment):
    request, result = assessment
    request = request.model_copy(
        update={
            "dpo_advice": "DPO ønsker dokumenteret evaluering af søgekvaliteten.",
            "data_subject_consultation": "Medarbejderrepræsentanter ønsker kortere slettefrist.",
        }
    )
    result.sections[0] = result.sections[0].model_copy(update={"source_ids": ["D1"]})
    result.risks[0] = result.risks[0].model_copy(
        update={
            "rationale": "Der mangler en ansvarlig for fejlmonitorering.",
            "consequences": "Registrerede risikerer forkert sagsbehandling.",
            "source_ids": ["D1"],
        }
    )
    result = result.model_copy(
        update={
            "version": 3,
            "case_db_id": "synthetic-case",
            "parent_assessment_id": "synthetic-docx-v2",
            "open_questions": ["Hvem gennemgår fejlloggen?"],
            "additional_risks": [
                {
                    "title": "Manglende logejer",
                    "scenario": "Fejl kan forblive uopdaget.",
                    "measures": "Udpeg en ansvarlig.",
                    "source_ids": ["D1"],
                }
            ],
            "ai_generation": {
                "model": "example/model",
                "prompt_version": "test-v1",
                "sources": [
                    {
                        "id": "D1",
                        "title": "Syntetisk procesbeskrivelse",
                        "document_id": "doc-1",
                        "source_url": "https://example.org/process",
                        "retrieved_at": "2026-09-20T09:15:00+00:00",
                        "api_key": "SECRET-DO-NOT-EXPORT",
                    },
                    {
                        "id": "D2",
                        "title": "Historisk kilde",
                        "url": "https://example.org/legacy",
                    },
                ],
                "limitations": ["Datagrundlaget er ufuldstændigt."],
            },
        }
    )
    doc = Document(BytesIO(export_dpia_docx(request, result)))
    text = text_content(doc)
    for value in [
        "Version: 3",
        "synthetic-case",
        "synthetic-docx-v2",
        "example/model",
        "test-v1",
        "Syntetisk procesbeskrivelse",
        "https://example.org/process",
        "Hentet: 2026-09-20T09:15:00+00:00",
        "https://example.org/legacy",
        "F1 Manglende logejer",
        "Hvem gennemgår fejlloggen?",
        request.dpo_advice,
        request.data_subject_consultation,
        "Der mangler en ansvarlig for fejlmonitorering.",
        "Registrerede risikerer forkert sagsbehandling.",
        "Kildehenvisninger: D1",
    ]:
        assert value in text
    assert "SECRET-DO-NOT-EXPORT" not in text
    assert "Disse forslag indgår ikke i de 33 skabelonrisici" in text
    assert doc.core_properties.identifier == result.id


def test_word_rejects_incomplete_or_inconsistent_snapshot(assessment):
    request, result = assessment
    with pytest.raises(TemplateError):
        export_dpia_docx(
            request, result.model_copy(update={"sections": result.sections[:-1]})
        )
    with pytest.raises(TemplateError):
        export_dpia_docx(
            request, result.model_copy(update={"risks": result.risks[:-1]})
        )
    risk = result.risks[0].model_copy(update={"inherent_risk": "low"})
    assert risk.likelihood == 2 and risk.impact == 3
    with pytest.raises(TemplateError):
        export_dpia_docx(
            request, result.model_copy(update={"risks": [risk] + result.risks[1:]})
        )


def test_word_escapes_plain_text_and_has_page_fields_without_wide_tables(assessment):
    request, result = assessment
    request = request.model_copy(update={"project_name": "Forslag <A&B>"})
    result.sections[0] = result.sections[0].model_copy(
        update={"text": "Bogstaveligt <tag> & tekst\x00."}
    )
    output = export_dpia_docx(request, result)
    doc = Document(BytesIO(output))
    assert "Bogstaveligt <tag> & tekst." in text_content(doc)
    assert all(len(t.columns) <= 5 for t in doc.tables)
    with ZipFile(BytesIO(output)) as archive:
        xml = archive.read("word/footer1.xml").decode()
        assert 'w:instr="PAGE"' in xml and 'w:instr="NUMPAGES"' in xml
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        assert (
            root.findall(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ins"
            )
            == []
        )


def test_word_describes_review_checks_without_claiming_legal_confidence(assessment):
    request, result = assessment
    result = result.model_copy(
        update={
            "summary_source_ids": ["D7"],
            "ai_generation": {
                "model": "example/model",
                "review": {
                    "status": "findings_require_review",
                    "rubric_version": "review-v2",
                    "threshold_note": "Kontrollen prioriterer menneskelig gennemgang.",
                    "checks": [
                        {
                            "id": "missing_evidence",
                            "label": "Manglende dokumentation",
                            "requires_review": True,
                            "probability": 0.97,
                            "section_ids": ["section:2.1", "risk:3.3"],
                        }
                    ],
                },
            },
        }
    )
    text = text_content(Document(BytesIO(export_dpia_docx(request, result))))
    assert "Manglende dokumentation: Kræver gennemgang" in text
    assert "review-v2" in text
    assert "Kontrollen prioriterer menneskelig gennemgang." in text
    assert "Berørte afsnit: section:2.1, risk:3.3" in text
    assert "Kildehenvisninger: D7" in text
    assert "0.97" not in text


def test_word_accepts_long_case_names_without_invalid_core_metadata(assessment):
    request, result = assessment
    long_name = "Langt projektnavn " * 25
    request = request.model_copy(
        update={"project_name": long_name, "organisation": "Organisation " * 30}
    )
    doc = Document(BytesIO(export_dpia_docx(request, result)))
    assert long_name in text_content(doc)
    assert len(doc.core_properties.title) == 255
    assert len(doc.core_properties.author) == 255
    assert len(doc.sections[0].header.paragraphs[0].text) < 130
    assert doc.styles["Title"].element.xpath(".//w:pBdr") == []


def test_word_formats_report_markdown_without_changing_snapshot_or_source_metadata(
    assessment,
):
    request, result = assessment
    result.executive_summary = (
        "## Dokumenteret\nAftalen beskriver **kommunens instrukser**.\n"
        "### Personoplysninger\nOplysningstyperne kræver kontrol.\n"
        "Skal afklares\n- **Hosting:** Den konkrete region mangler.\n"
        "- Sletning er ikke dokumenteret.\n"
        "Før godkendelse\n3. Indhent leverandørens svar.\n7. Få en faglig beslutning."
    )
    result.sections[0] = result.sections[0].model_copy(
        update={
            "text": "## Sikkerhed\nSkal afklares: **Adgangskontrollen** skal verificeres.\n<script>ingen handling</script>",
        }
    )
    result.ai_generation = {
        "sources": [
            {
                "id": "synthetic-source",
                "title": "## Kildetitel med **markører**",
            }
        ]
    }
    before = deepcopy(result.model_dump(mode="json"))
    word = Document(BytesIO(export_dpia_docx(request, result)))
    heading = next(p for p in word.paragraphs if p.text == "Dokumenteret")
    subheading = next(p for p in word.paragraphs if p.text == "Personoplysninger")
    assert heading.style.name == "Heading 3"
    assert subheading.style.name == "Heading 4"
    assert (
        next(p for p in word.paragraphs if p.text == "Skal afklares").style.name
        == "Heading 3"
    )
    statement = next(
        p
        for p in word.paragraphs
        if p.text == "Aftalen beskriver kommunens instrukser."
    )
    assert any(
        run.text == "kommunens instrukser" and run.bold for run in statement.runs
    )
    bullet = next(
        p for p in word.paragraphs if p.text == "Hosting: Den konkrete region mangler."
    )
    assert bullet.style.name == "List Bullet"
    assert any(run.text == "Hosting:" and run.bold for run in bullet.runs)
    text = text_content(word)
    for preserved in (
        "3. Indhent leverandørens svar.",
        "7. Få en faglig beslutning.",
        "Sletning er ikke dokumenteret.",
        "Skal afklares: Adgangskontrollen skal verificeres.",
        "<script>ingen handling</script>",
        "## Kildetitel med **markører**",
    ):
        assert preserved in text
    assert "## Dokumenteret" not in text and "**Adgangskontrollen**" not in text
    assert result.model_dump(mode="json") == before


def test_word_report_formatting_preserves_literal_markup_and_safe_empty_fragments():
    from src.services.dpia_docx import _structured_text

    word = Document()
    _structured_text(word, "", keep_with_next=True)
    assert not word.paragraphs
    _structured_text(
        word,
        "Bogstaveligt \\**ikke fed**\nKilden hedder vendor__data__id.\n`**rå tekst**`\n**Uafsluttet",
        heading_level=4,
    )
    assert [p.text for p in word.paragraphs] == [
        "Bogstaveligt \\**ikke fed**",
        "Kilden hedder vendor__data__id.",
        "**rå tekst**",
        "**Uafsluttet",
    ]
