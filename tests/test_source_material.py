"""Isolated supplier intake: real formats, traceable excerpts, atomic storage and SSRF."""

from datetime import UTC, datetime
import hashlib
from io import BytesIO
import socket
import zipfile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from src.api.source_material import router
from src.api.workspace import WORKSPACE_ACCESS, router as workspace_router
from src.auth import UserPrincipal, get_current_user
from src.database.cases import create_case
from src.database.connection import Base, get_db
from src.database.document_bank import (
    MunicipalDocument,
    MunicipalDocumentVersion,
    CaseDocumentLink,
    link_document_to_case,
)
from src.services import source_material as material
from src.services import safe_public_fetch as public
from src.services.document_bank_storage import validate_office_package


def test_standard_pptx_printer_settings_are_not_treated_as_macros():
    raw = pptx_bytes(
        extra={"ppt/printerSettings/printerSettings1.bin": b"inert printer settings"}
    )
    buffer = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as original, zipfile.ZipFile(
        buffer, "w"
    ) as updated:
        for name in original.namelist():
            content = original.read(name)
            if name == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.presentationml.printerSettings"/></Types>',
                )
            updated.writestr(name, content)
    validate_office_package(".pptx", buffer.getvalue())
    with pytest.raises(ValueError, match="printer settings"):
        validate_office_package(".pptx", raw)


def pptx_bytes(*, extra=None, reverse=False):
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/></Types>',
        )
        order = [2, 1] if reverse else [1, 2]
        archive.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst>'
            + "".join(
                f'<p:sldId id="{number}" r:id="rId{number}"/>' for number in order
            )
            + "</p:sldIdLst></p:presentation>",
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="slides/slide1.xml"/><Relationship Id="rId2" Target="slides/slide2.xml"/></Relationships>',
        )
        archive.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><a:p><a:r><a:t>Kommunalt fagsystem</a:t></a:r></a:p></p:sld>',
        )
        archive.writestr(
            "ppt/slides/slide2.xml",
            f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><a:tbl><a:tr><a:tc><a:p><a:r><a:t>Databehandler</a:t></a:r></a:p></a:tc><a:tc><a:p><a:r><a:t>Leverandørens oplysning</a:t></a:r></a:p></a:tc></a:tr></a:tbl></p:sld>',
        )
        for name, content in (extra or {}).items():
            archive.writestr(name, content)
    return output.getvalue()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    storage = tmp_path / "files"
    monkeypatch.setenv("DOCUMENT_BANK_STORAGE_DIR", str(storage))
    engine = create_engine(f"sqlite:///{tmp_path / 'cases.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        first = create_case(
            db, case_id="SKOLE-1", title="Indkøb af kommunalt fagsystem"
        )
        second = create_case(db, case_id="SUNDHED-2", title="Et andet fagsystem")
        case_ids = (first.id, second.id)
        db.commit()
    principal = UserPrincipal(
        oid="test-user",
        name="Kommunal sagsbehandler",
        roles=["Hammeren.Sagsbehandler"],
        auth_mode="development",
        identity_assurance="development_only",
    )
    app = FastAPI()
    app.include_router(router)
    app.include_router(workspace_router)

    def db_override():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = db_override
    app.dependency_overrides[WORKSPACE_ACCESS] = lambda: principal
    with TestClient(app) as client:
        yield client, factory, case_ids, storage, app, principal
    engine.dispose()


def test_powerpoint_upload_keeps_real_slide_order_tables_original_and_unapproved_state(
    setup,
):
    client, factory, (case_id, other), _, _, _ = setup
    content = pptx_bytes(reverse=True)
    response = client.post(
        f"/api/v3/cases/{case_id}/source-material",
        files={"file": ("systemgennemgang.pptx", content)},
        data={"title": "Leverandørens systemgennemgang"},
    )
    assert response.status_code == 201, response.text
    item = response.json()
    assert item["status"] == "extracted"
    assert item["document_status"] == "draft"
    assert item["review_status"] == "unreviewed"
    assert item["checksum"] == hashlib.sha256(content).hexdigest()
    assert item["excerpts"][0]["locator"] == "Slide 1"
    assert "Leverandørens oplysning" in item["excerpts"][0]["text"]
    assert item["excerpts"][1]["locator"] == "Slide 2"
    assert item["excerpts"][1]["text"] == "Kommunalt fagsystem"
    assert client.get(item["download_url"]).content == content
    assert client.get(f"/api/v3/cases/{other}/source-material").json() == {
        "items": [],
        "count": 0,
        "analysis_limits": {
            "max_documents": 25,
            "max_total_text_chars": 500_000,
            "max_document_text_chars": 200_000,
            "max_total_excerpts": 1_000,
            "max_document_excerpts": 500,
            "max_file_bytes": 5_000_000,
            "batching_enabled": True,
            "max_batches": 20,
            "max_case_documents": 100,
            "max_case_text_chars": 5_000_000,
            "max_case_excerpts": 10_000,
            "max_extraction_text_chars": 2_000_000,
            "max_extraction_segments": 5_000,
        },
    }
    with factory() as db:
        version = db.get(MunicipalDocumentVersion, item["version_id"])
        assert version.approved_at is None and version.approved_by is None
        evidence = material.case_source_evidence(db, case_id)
        assert len(evidence) == 2
        assert evidence[0]["id"] == f"document:{version.id}:1"
        assert evidence[0]["checksum"] == item["checksum"]
        assert material.case_source_evidence(db, other) == []


def test_word_paragraphs_and_table_rows_are_preserved_with_locators():
    doc = Document()
    doc.add_paragraph("Formål: sagsbehandling på sundhedsområdet.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Oplysninger"
    table.cell(0, 1).text = "Behandlingssted"
    table.cell(1, 0).text = "Helbredsoplysninger"
    table.cell(1, 1).text = "Leverandøren oplyser EU"
    doc.add_paragraph("Oplysningen er ikke efterprøvet.")
    output = BytesIO()
    doc.save(output)
    extracted = material.extract_source(output.getvalue(), "databehandleraftale.docx")
    assert [item["locator"] for item in extracted.excerpts] == [
        "Afsnit 1",
        "Tabel 1, række 1",
        "Tabel 1, række 2",
        "Afsnit 2",
    ]
    assert (
        extracted.excerpts[2]["text"] == "Helbredsoplysninger | Leverandøren oplyser EU"
    )


def test_pdf_page_text_keeps_page_locator():
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    for label in ["Data controller", "Data processor"]:
        writer.add_blank_page(width=500, height=500)
        page = writer.pages[-1]
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 450 Td ({label}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
    output = BytesIO()
    writer.write(output)
    result = material.extract_source(output.getvalue(), "processor.pdf")
    assert [(item["locator"], item["text"]) for item in result.excerpts] == [
        ("Side 1", "Data controller"),
        ("Side 2", "Data processor"),
    ]


def test_text_and_segment_limits_are_disclosed():
    result = material.extract_source(("x" * 2_000_002).encode(), "lang.txt")
    assert sum(len(item["text"]) for item in result.excerpts) == 2_000_000
    assert not result.complete
    assert any("afkortet" in warning for warning in result.warnings)
    result = material.extract_source(
        "\n\n".join(str(index) for index in range(5005)).encode(), "lang.txt"
    )
    assert len(result.excerpts) == 5000
    assert not result.complete
    assert any("afkortet" in warning for warning in result.warnings)


@pytest.mark.parametrize(
    "extra",
    [
        {"ppt/vbaProject.bin": b"macro"},
        {"../outside.xml": b"<root/>"},
        {"ppt/unsafe.xml": b'<!DOCTYPE r [<!ENTITY x "injected">]><r>&x;</r>'},
        {"ppt/bomb.txt": b"x" * 400_000},
    ],
)
def test_malicious_office_packages_are_rejected(extra):
    with pytest.raises(ValueError):
        validate_office_package(".pptx", pptx_bytes(extra=extra))


def test_fake_office_signature_and_legacy_ppt_and_oversized_upload_are_rejected(setup):
    client, _, (case_id, _), _, _, _ = setup
    for filename, content, expected in [
        ("old.ppt", b"legacy", "pptx"),
        ("fake.pptx", b"PK-not-office", "ZIP"),
        ("large.txt", b"x" * 5_000_001, "5 MB"),
    ]:
        response = client.post(
            f"/api/v3/cases/{case_id}/source-material",
            files={"file": (filename, content)},
        )
        assert response.status_code in {413, 422}
        assert expected in response.json()["detail"]


def test_link_failure_rolls_back_database_and_exact_storage_object(setup, monkeypatch):
    _, factory, (case_id, _), storage, _, _ = setup

    def fail(*args, **kwargs):
        raise RuntimeError("synthetic link failure")

    monkeypatch.setattr(material, "link_document_to_case", fail)
    with factory() as db:
        with pytest.raises(RuntimeError):
            material.save_case_source(
                db,
                case_id,
                filename="grundlag.txt",
                content=b"Municipal evidence",
                actor="Sagsbehandler",
            )
        assert db.query(MunicipalDocument).count() == 0
        assert db.query(MunicipalDocumentVersion).count() == 0
        assert db.query(CaseDocumentLink).count() == 0
    assert not list(storage.rglob("*.txt"))


def test_unknown_case_is_404_and_has_no_storage_side_effects(setup):
    client, factory, _, storage, _, _ = setup
    response = client.post(
        "/api/v3/cases/missing/source-material",
        files={"file": ("grundlag.txt", b"Municipal evidence")},
    )
    assert response.status_code == 404
    with factory() as db:
        assert db.query(MunicipalDocument).count() == 0
    assert not storage.exists()


def test_corrupted_source_and_output_cannot_be_used_as_evidence(setup):
    client, factory, (case_id, _), storage, _, _ = setup
    saved = client.post(
        f"/api/v3/cases/{case_id}/source-material",
        files={"file": ("grundlag.txt", b"Supplier says EU")},
    ).json()
    with factory() as db:
        version = db.get(MunicipalDocumentVersion, saved["version_id"])
        path = storage / version.storage_key
        path.write_bytes(b"tampered")
        assert material.case_source_evidence(db, case_id) == []
        assert (
            material.list_case_source_material(db, case_id)[0]["status"] == "unreadable"
        )
        path.write_bytes(b"Supplier says EU")
        link_document_to_case(
            db,
            case_db_id=case_id,
            document_id=version.document_id,
            document_version_id=version.id,
            link_role="output",
            allow_unapproved=True,
        )
        assert material.case_source_evidence(db, case_id) == []


def public_dns(address="93.184.216.34"):
    return [
        (
            socket.AF_INET6 if ":" in address else socket.AF_INET,
            socket.SOCK_STREAM,
            6,
            "",
            (address, 443),
        )
    ]


@pytest.mark.parametrize(
    "address",
    [
        "127.0.0.1",
        "10.0.0.8",
        "192.168.1.2",
        "169.254.169.254",
        "100.64.1.2",
        "0.0.0.0",
        "224.0.0.1",
        "::1",
        "fd00::1",
        "fe80::1",
        "::ffff:127.0.0.1",
        "2002:7f00:1::",
    ],
)
def test_public_fetch_rejects_private_reserved_and_tunnel_addresses(
    monkeypatch, address
):
    monkeypatch.setattr(
        public.socket, "getaddrinfo", lambda *args, **kwargs: public_dns(address)
    )
    with pytest.raises(public.PublicSourceError, match="offentlig hjemmeside"):
        public.validate_public_url("https://supplier.example/")


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "ftp://supplier.example",
        "https://user:pass@supplier.example",
        "https://supplier.example:443/",
        "https://supplier.example\\@localhost/",
        "https://[fe80::1%25en0]/",
    ],
)
def test_unsafe_url_syntax_is_rejected_before_dns(monkeypatch, url):
    monkeypatch.setattr(
        public.socket,
        "getaddrinfo",
        lambda *args, **kwargs: pytest.fail("DNS must not be called"),
    )
    with pytest.raises(public.PublicSourceError):
        public.validate_public_url(url)


class FakeResponse:
    def __init__(self, status=200, headers=None, content=b"<p>Fagsystem</p>"):
        self.status = status
        self.headers = headers or {"Content-Type": "text/html"}
        self.body = BytesIO(content)

    def getheader(self, name, default=None):
        return self.headers.get(name, default)

    def read1(self, size):
        return self.body.read(size)


def test_every_redirect_is_checked_and_private_target_is_never_connected(monkeypatch):
    connected = []
    monkeypatch.setattr(
        public.socket,
        "getaddrinfo",
        lambda host, *args, **kwargs: public_dns(
            "127.0.0.1" if host == "internal.example" else "93.184.216.34"
        ),
    )

    class Connection:
        sock = None

        def __init__(self, host, port, ip, timeout):
            connected.append((host, ip))

        def request(self, *args, **kwargs):
            pass

        def getresponse(self):
            return FakeResponse(302, {"Location": "https://internal.example/private"})

        def close(self):
            pass

    monkeypatch.setattr(public, "_PinnedHTTPSConnection", Connection)
    with pytest.raises(public.PublicSourceError, match="offentlig hjemmeside"):
        public.fetch_public_source("https://supplier.example/")
    assert connected == [("supplier.example", "93.184.216.34")]


def test_connection_uses_validated_ip_and_tls_original_hostname(monkeypatch):
    sockets = []
    wraps = []

    class Context:
        def wrap_socket(self, sock, *, server_hostname):
            wraps.append(server_hostname)
            return sock

    monkeypatch.setattr(
        public.socket,
        "create_connection",
        lambda destination, timeout: sockets.append(destination) or object(),
    )
    monkeypatch.setattr(public.ssl, "create_default_context", lambda: Context())
    connection = public._PinnedHTTPSConnection(
        "supplier.example", 443, "93.184.216.34", 2
    )
    connection.connect()
    assert sockets == [("93.184.216.34", 443)]
    assert wraps == ["supplier.example"]


def test_url_html_snapshot_is_script_free_immutable_and_traceable(setup, monkeypatch):
    client, factory, (case_id, _), _, _, _ = setup
    raw = "<html><head><title>Title</title><script>unsafe()</script></head><body><h1>Kommunalt fagsystem</h1><p>Leverandøren oplyser <strong>EU</strong>.</p><script>send_credentials()</script></body></html>".encode()
    fetched = public.PublicSource(
        "https://supplier.example/product",
        raw,
        "text/html",
        datetime(2026, 9, 20, tzinfo=UTC).isoformat(),
        hashlib.sha256(raw).hexdigest(),
    )
    monkeypatch.setattr(material, "fetch_public_source", lambda url: fetched)
    response = client.post(
        f"/api/v3/cases/{case_id}/source-material/url",
        json={"url": fetched.url, "title": "Leverandørens produktbeskrivelse"},
    )
    assert response.status_code == 201, response.text
    item = response.json()
    assert item["source_url"] == fetched.url
    assert item["retrieved_at"] == fetched.retrieved_at
    text = client.get(item["download_url"]).content.decode()
    assert text == "Kommunalt fagsystem\nLeverandøren oplyser EU."
    with factory() as db:
        version = db.get(MunicipalDocumentVersion, item["version_id"])
        assert (
            version.version_metadata["raw_source_sha256"]
            == hashlib.sha256(raw).hexdigest()
        )
        assert version.version_metadata["text_snapshot_format"] == "html-blocks-2"
        evidence = material.case_source_evidence(db, case_id)
        assert len(evidence) == 1
        assert evidence[0]["text"] == text
        assert evidence[0]["source_url"] == fetched.url


def test_html_microparagraphs_keep_all_text_and_order_within_excerpt_budget():
    paragraphs = [f"FAQ {number}: Opbevaring i EU." for number in range(650)]
    raw = (
        "<html><body>"
        + "".join(
            f"<div>FAQ {number}: Opbevaring <strong>i</strong> EU.</div>"
            for number in range(650)
        )
        + "<p>Sidste svar: gennemgå <em>alle</em> vilkår.</p></body></html>"
    ).encode()
    paragraphs.append("Sidste svar: gennemgå alle vilkår.")
    source = public.PublicSource(
        "https://supplier.example/faq",
        raw,
        "text/html",
        datetime(2026, 9, 21, tzinfo=UTC).isoformat(),
        hashlib.sha256(raw).hexdigest(),
    )
    filename, content, _ = material.web_snapshot(source)
    extracted = material.extract_source(content, filename)
    assert len(extracted.excerpts) < 200
    assert all(len(item["text"]) <= 2_000 for item in extracted.excerpts)
    assert " ".join(content.decode().split()) == " ".join(paragraphs)
    assert " ".join(
        " ".join(item["text"] for item in extracted.excerpts).split()
    ) == " ".join(paragraphs)
    assert not any("afkortet" in warning for warning in extracted.warnings)


def test_html_long_paragraphs_keep_words_urls_and_existing_character_limit():
    long_url = "https://supplier.example/" + "x" * 2_000
    paragraph = " ".join(f"Betingelse-{number}" for number in range(20_000))
    raw = f"<p>{long_url}</p><p>{paragraph}</p><p>Sidste vilkår</p>".encode()
    source = public.PublicSource(
        "https://supplier.example/privacy",
        raw,
        "text/html",
        datetime(2026, 9, 21, tzinfo=UTC).isoformat(),
        hashlib.sha256(raw).hexdigest(),
    )
    filename, content, _ = material.web_snapshot(source)
    assert " ".join(content.decode().split()) == (
        f"{long_url} {paragraph} Sidste vilkår"
    )
    assert long_url in content.decode().split("\n\n")
    extracted = material.extract_source(content, filename)
    assert sum(len(item["text"]) for item in extracted.excerpts) > 200_000
    assert len(extracted.excerpts) < 200
    assert extracted.complete
    assert not any("afkortet" in warning for warning in extracted.warnings)


def test_historical_text_snapshot_keeps_bytes_excerpt_ids_and_boundaries(setup):
    client, factory, (case_id, _), _, _, _ = setup
    original = b"Historical heading\n\nFirst supplier statement.\n\nLast answer."
    with factory() as db:
        saved = material.save_case_source(
            db,
            case_id,
            filename="supplier.example.txt",
            content=original,
            actor="test-user",
            provenance={"source_url": "https://supplier.example/old"},
        )
    assert client.get(saved["download_url"]).content == original
    listed = client.get(f"/api/v3/cases/{case_id}/source-material").json()["items"]
    assert listed[0]["excerpts"] == [
        {
            "id": f"document:{saved['version_id']}:{number}",
            "locator": f"Afsnit {number}",
            "text": text,
        }
        for number, text in enumerate(original.decode().split("\n\n"), 1)
    ]
    with factory() as db:
        version = db.get(MunicipalDocumentVersion, saved["version_id"])
        assert "text_snapshot_format" not in version.version_metadata
        evidence = material.case_source_evidence(db, case_id)
        assert [item["id"] for item in evidence] == [
            item["id"] for item in listed[0]["excerpts"]
        ]


def test_plain_text_web_snapshot_keeps_original_paragraph_boundaries():
    original = b"Heading\n\nFirst statement.\n\nLast answer."
    source = public.PublicSource(
        "https://supplier.example/policy.txt",
        original,
        "text/plain",
        datetime(2026, 9, 21, tzinfo=UTC).isoformat(),
        hashlib.sha256(original).hexdigest(),
    )
    filename, content, metadata = material.web_snapshot(source)
    assert content == original
    assert "text_snapshot_format" not in metadata
    assert len(material.extract_source(content, filename).excerpts) == 3


def test_upload_and_url_routes_require_workspace_role(setup):
    client, _, (case_id, _), _, app, principal = setup
    app.dependency_overrides.pop(WORKSPACE_ACCESS)
    principal.roles = []
    app.dependency_overrides[get_current_user] = lambda: principal
    for method, route in [("GET", ""), ("POST", ""), ("POST", "/url")]:
        response = client.request(
            method,
            f"/api/v3/cases/{case_id}/source-material{route}",
            json={"url": "https://example.com/"},
        )
        assert response.status_code == 403


def test_client_supplied_cached_excerpts_cannot_override_original_evidence(setup):
    client, factory, (case_id, _), _, _, _ = setup
    saved = client.post(
        f"/api/v3/cases/{case_id}/source-material",
        files={"file": ("grundlag.txt", b"Supplier statement, not a legal decision")},
    ).json()
    with factory() as db:
        version = db.get(MunicipalDocumentVersion, saved["version_id"])
        version.version_metadata = {
            "source_material": {
                "excerpts": [{"locator": "Afsnit 1", "text": "FORGED APPROVAL"}]
            }
        }
        db.commit()
        evidence = material.case_source_evidence(db, case_id)
        assert evidence[0]["text"] == "Supplier statement, not a legal decision"


def test_mixed_public_private_dns_response_is_rejected(monkeypatch):
    monkeypatch.setattr(
        public.socket,
        "getaddrinfo",
        lambda *args, **kwargs: public_dns() + public_dns("10.0.0.1"),
    )
    with pytest.raises(public.PublicSourceError):
        public.validate_public_url("https://supplier.example/")


@pytest.mark.parametrize(
    "headers,content",
    [
        ({"Content-Type": "text/html", "Content-Length": "5000001"}, b"x"),
        ({"Content-Type": "text/html"}, b"x" * 5_000_001),
        ({"Content-Type": "text/html", "Content-Encoding": "gzip"}, b"x"),
        ({"Content-Type": "application/octet-stream"}, b"executable"),
    ],
)
def test_public_fetch_rejects_unbounded_and_unsupported_responses(
    monkeypatch, headers, content
):
    monkeypatch.setattr(
        public.socket, "getaddrinfo", lambda *args, **kwargs: public_dns()
    )

    class Connection:
        sock = None

        def __init__(self, *args):
            pass

        def request(self, *args, **kwargs):
            assert "Cookie" not in kwargs["headers"]
            assert "Authorization" not in kwargs["headers"]

        def getresponse(self):
            return FakeResponse(headers=headers, content=content)

        def close(self):
            pass

    monkeypatch.setattr(public, "_PinnedHTTPSConnection", Connection)
    with pytest.raises(public.PublicSourceError):
        public.fetch_public_source("https://supplier.example/")


def test_office_rejects_duplicate_entries_encryption_and_wrong_signature():
    source = pptx_bytes()
    output = BytesIO(source)
    with pytest.warns(UserWarning), zipfile.ZipFile(output, "a") as archive:
        archive.writestr("ppt/slides/slide1.xml", "<slide/>")
    with pytest.raises(ValueError, match="duplicate"):
        validate_office_package(".pptx", output.getvalue())
    encrypted = bytearray(source)
    start = encrypted.index(b"PK\x01\x02")
    encrypted[start + 8] |= 1
    with pytest.raises(ValueError, match="Encrypted"):
        validate_office_package(".pptx", bytes(encrypted))
    with pytest.raises(ValueError, match="match"):
        validate_office_package(".docx", source)


def test_large_paragraph_keeps_every_character_in_stable_parts(setup):
    _, factory, (case_id, _), _, _, _ = setup
    text = "A" * 200_000 + "B" * 200_000 + "LAST EVIDENCE"
    with factory() as db:
        saved = material.save_case_source(
            db, case_id, filename="large.txt", content=text.encode(), actor="Tester"
        )
        assert saved["extraction_complete"]
        sources = material.case_source_evidence(db, case_id)
        prefix = f"document:{saved['version_id']}:1"
        assert [source["id"] for source in sources] == [
            prefix,
            prefix + ":part:2",
            prefix + ":part:3",
        ]
        assert "".join(source["text"] for source in sources) == text
        assert sources[-1]["locator"] == "Afsnit 1, tegn 400001–400013"


def test_incomplete_extraction_cannot_become_a_partial_analysis(setup):
    _, factory, (case_id, _), _, _, _ = setup
    with factory() as db:
        saved = material.save_case_source(
            db,
            case_id,
            filename="oversized.txt",
            content=b"x" * 2_000_001,
            actor="Tester",
        )
        assert not saved["extraction_complete"]
        with pytest.raises(ValueError, match="Ingen delvis analyse"):
            material.case_source_evidence(db, case_id)
