"""Report editing preserves immutable legal/risk evidence and honest provenance."""

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
from threading import Barrier
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import create_engine
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker

from src.api import dpia_revisions as api
from src.auth import UserPrincipal
from src.database.cases import Case, create_case
from src.database.connection import Base, get_db
from src.database.dpia import (
    DPIAAssessmentRecord,
    assessment_result_payload,
    save_assessment,
)
from src.services import dpia_revisions as service
from src.services.dpia_assessment import DPIAAssessmentResponse
from src.services.dpia_docx import export_dpia_docx
from src.services.dpia_export import export_dpia_xlsx
from tests.test_dpia_assessment import make_assessment


@pytest.fixture
def setup(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'revisions.db'}", connect_args={"timeout": 10}
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    request, result = make_assessment()
    editable = next(
        section
        for section in result.sections
        if section.review_status == "requires_review"
    )
    result.ai_generation = {
        "model": "synthetic-model",
        "review": {
            "model": "synthetic-jev",
            "checks": [
                {
                    "id": f"section:{editable.id}",
                    "label": "Kildekontrol",
                    "section_ids": [f"section:{editable.id}"],
                    "requires_review": False,
                },
                {
                    "id": "summary",
                    "label": "Resumé",
                    "section_ids": ["summary"],
                    "requires_review": False,
                },
            ],
        },
        "sources": [
            {"id": "input:purpose", "title": "Sagens formål", "text": request.purpose}
        ],
    }
    editable.source_ids = ["input:purpose"]
    with factory() as db:
        case = create_case(db, case_id="FS-TEST-REVISION", title=request.project_name)
        base = save_assessment(
            db,
            assessment_id=result.id,
            created_at=result.created_at,
            request_payload=request.model_dump(mode="json"),
            result_payload=result.model_dump(mode="json"),
            case_db_id=case.id,
        )
        base_snapshot = deepcopy(base.result_payload)
        db.commit()
        db.refresh(case)
        case_snapshot = case.to_dict()
    app = FastAPI()
    app.include_router(api.router)

    def database():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = database
    app.dependency_overrides[api.WORKSPACE_ACCESS] = lambda: UserPrincipal(
        oid="reviewer-1",
        name="Faglig reviewer",
        roles=["Hammeren.Sagsbehandler"],
        auth_mode="development",
        identity_assurance="development_only",
    )
    yield {
        "client": TestClient(app),
        "factory": factory,
        "base": base_snapshot,
        "request": request,
        "section": editable.id,
        "case": case_snapshot,
        "app": app,
    }
    engine.dispose()


def payload(setup, **overrides):
    data = {
        "changes": [
            {
                "kind": "section",
                "target_id": setup["section"],
                "field": "text",
                "text": "Kommunens faglige gennemgang præciserer anvendelsen af sagens oplysninger.",
                "source_ids": ["input:purpose"],
            }
        ],
        "note": "Præcisering aftalt ved den faglige gennemgang.",
        "request_id": str(uuid4()),
    }
    return {**data, **overrides}


def endpoint(setup, base_id=None):
    return f"/api/dpia/assessments/{base_id or setup['base']['id']}/revisions"


def test_edit_persists_a_new_version_without_changing_risk_or_legal_decisions(setup):
    base = setup["base"]
    changed = payload(setup)
    changed["changes"].append(
        {
            "kind": "risk",
            "target_id": base["risks"][0]["id"],
            "field": "measures",
            "text": "En dokumenteret adgangstest skal gennemføres før ibrugtagning.",
            "source_ids": ["input:purpose"],
        }
    )
    response = setup["client"].post(endpoint(setup), json=changed)
    assert response.status_code == 201, response.text
    saved = response.json()
    assert saved["id"] != base["id"] and saved["version"] == 2
    assert saved["parent_assessment_id"] == base["id"]
    for key in (
        "status",
        "status_label",
        "risk_level",
        "completeness",
        "dpia_required",
        "screening_criteria",
        "legal_verification",
        "missing_information",
        "blockers",
        "next_steps",
        "ai_generation",
    ):
        assert saved[key] == base[key]
    for original, updated in zip(base["risks"], saved["risks"]):
        for key in (
            "id",
            "likelihood",
            "impact",
            "inherent_risk",
            "residual_likelihood",
            "residual_impact",
            "residual_risk",
            "owner",
            "implementation_status",
            "due_date",
        ):
            assert updated[key] == original[key]
    metadata = saved["editorial_revision"]
    assert metadata["edited_by"] == "Faglig reviewer"
    assert metadata["actor_id"] == "reviewer-1"
    assert metadata["stale_check_ids"] == sorted(
        [f"section:{setup['section']}", f"risk:{base['risks'][0]['id']}"]
    )
    assert metadata["review_status"] == "pending_recheck"
    assert metadata["changes"][0]["after"]["text"] == changed["changes"][0]["text"]
    with setup["factory"]() as db:
        assert db.get(DPIAAssessmentRecord, base["id"]).result_payload == base
        assert db.get(DPIAAssessmentRecord, saved["id"]).request_payload == setup[
            "request"
        ].model_dump(mode="json")
        assert db.get(Case, base["case_db_id"]).to_dict() == setup["case"]


def test_retry_is_idempotent_but_other_edits_and_stale_bases_are_rejected(setup):
    body = payload(setup)
    first = setup["client"].post(endpoint(setup), json=body)
    assert first.status_code == 201
    assert setup["client"].post(endpoint(setup), json=body).json() == first.json()
    assert setup["client"].post(endpoint(setup), json=payload(setup)).status_code == 409
    body["note"] = "En anden begrundelse for en anden gemmehandling."
    assert setup["client"].post(endpoint(setup), json=body).status_code == 409
    history = setup["client"].get(endpoint(setup)).json()
    assert history["latest_id"] == first.json()["id"]
    assert [item["version"] for item in history["items"]] == [2, 1]


def test_later_human_edits_keep_previous_targets_pending_and_history_scoped_to_case(
    setup,
):
    first = setup["client"].post(endpoint(setup), json=payload(setup)).json()
    second_payload = payload(
        setup,
        changes=[
            {
                "kind": "scope",
                "target_id": "scope",
                "field": "scope",
                "text": "Afgrænsningen vedrører alene kommunens beskrevne administrative anvendelse.",
                "source_ids": [],
            }
        ],
    )
    second = setup["client"].post(endpoint(setup, first["id"]), json=second_payload)
    assert second.status_code == 201, second.text
    assert second.json()["editorial_revision"]["stale_check_ids"] == sorted(
        ["summary", f"section:{setup['section']}"]
    )
    with setup["factory"]() as db:
        other = create_case(db, case_id="FS-OTHER", title="En anden sag")
        other_result = deepcopy(setup["base"])
        save_assessment(
            db,
            assessment_id=str(uuid4()),
            created_at=datetime.now(UTC),
            request_payload=setup["request"].model_dump(mode="json"),
            result_payload=other_result,
            case_db_id=other.id,
        )
        db.commit()
    assert len(setup["client"].get(endpoint(setup)).json()["items"]) == 3


@pytest.mark.parametrize("kind", ["section", "risk", "summary", "scope"])
def test_prose_changes_require_rechecking_advice_but_preserve_its_history(setup, kind):
    recommendation = {
        "id": "local_processing",
        "title": "Overvej isoleret behandling",
        "proposal": "Afprøv lokal behandling, hvis løsningen understøtter den.",
        "rationale": "Mødereferater kan indeholde fortrolige oplysninger.",
        "prerequisites": "Afklar driftsansvar og teknisk mulighed.",
        "verification": "Dokumentér en afprøvning af kvalitet, adgang og sletning.",
        "source_ids": ["input:purpose"],
    }
    check = {
        "id": "recommendation:local_processing",
        "label": "Forslag om isoleret behandling",
        "section_ids": ["recommendation:local_processing"],
        "probability": 0.2,
        "requires_review": False,
    }
    with setup["factory"]() as db:
        base = db.get(DPIAAssessmentRecord, setup["base"]["id"])
        snapshot = deepcopy(base.result_payload)
        snapshot["recommendations"] = [recommendation]
        snapshot["ai_generation"]["review"]["checks"].append(check)
        base.result_payload = snapshot
        db.commit()
    target_id, field = {
        "section": (setup["section"], "text"),
        "risk": (setup["base"]["risks"][0]["id"], "scenario"),
        "summary": ("summary", "executive_summary"),
        "scope": ("scope", "scope"),
    }[kind]
    response = setup["client"].post(endpoint(setup), json=payload(setup, changes=[{
        "kind": kind, "target_id": target_id, "field": field,
        "text": "Anvendelsen afgrænses til administrative møder uden borgeroplysninger.",
        "source_ids": ["input:purpose"],
    }]))
    assert response.status_code == 201, response.text
    saved = response.json()
    assert "recommendation:local_processing" in saved["editorial_revision"]["stale_check_ids"]
    assert saved["recommendations"] == [recommendation]
    assert saved["ai_generation"] == snapshot["ai_generation"]
    assert saved["status"] == snapshot["status"]
    assert saved["blockers"] == snapshot["blockers"]
    with setup["factory"]() as db:
        assert db.get(DPIAAssessmentRecord, setup["base"]["id"]).result_payload == snapshot


@pytest.mark.parametrize(
    "mutation",
    [
        lambda body: body.update(status="ready_for_review"),
        lambda body: body["changes"][0].update(field="likelihood", text="1"),
        lambda body: body["changes"][0].update(source_ids=["document:other-case"]),
        lambda body: body["changes"][0].update(target_id="unknown"),
        lambda body: body["changes"][0].update(
            source_ids=["input:purpose", "input:purpose"]
        ),
        lambda body: body.update(changes=body["changes"] * 2),
        lambda body: body.update(note=""),
        lambda body: body["changes"][0].update(text=" " * 10),
        lambda body: body["changes"][0].update(text="x" * 20_001),
    ],
)
def test_invalid_or_overreaching_edits_do_not_create_versions(setup, mutation):
    body = payload(setup)
    mutation(body)
    response = setup["client"].post(endpoint(setup), json=body)
    assert response.status_code == 422, response.text
    assert len(setup["client"].get(endpoint(setup)).json()["items"]) == 1


def test_missing_information_and_not_applicable_sections_cannot_be_written_away(setup):
    protected = [
        item
        for item in setup["base"]["sections"]
        if item["review_status"] != "requires_review"
    ]
    assert protected
    for section in protected:
        body = payload(setup)
        body["changes"][0]["target_id"] = section["id"]
        assert setup["client"].post(endpoint(setup), json=body).status_code == 422


def test_failed_save_rolls_back_all_writes(setup, monkeypatch):
    real_save = service.save_assessment

    def fail_after_insert(*args, **kwargs):
        real_save(*args, **kwargs)
        raise SQLAlchemyError("Synthetic failure")

    monkeypatch.setattr(service, "save_assessment", fail_after_insert)
    assert setup["client"].post(endpoint(setup), json=payload(setup)).status_code == 503
    assert len(setup["client"].get(endpoint(setup)).json()["items"]) == 1


def test_concurrent_revisions_never_overwrite_or_fork_the_latest_base(setup):
    barrier = Barrier(2)

    def submit():
        with setup["factory"]() as db:
            barrier.wait()
            try:
                record = service.revise_assessment(
                    db,
                    setup["base"]["id"],
                    service.RevisionInput.model_validate(payload(setup)),
                    actor_id="reviewer",
                    actor_name="Reviewer",
                )
                db.commit()
                return record.version
            except service.RevisionConflict:
                db.rollback()
                return "conflict"

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: submit(), range(2)))
    assert sorted(str(value) for value in results) == ["2", "conflict"]


def test_revised_word_and_excel_disclose_edit_and_stale_jev_check(setup):
    saved = setup["client"].post(endpoint(setup), json=payload(setup)).json()
    result = DPIAAssessmentResponse.model_validate(saved)
    doc = Document(BytesIO(export_dpia_docx(setup["request"], result)))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Fagligt redigeret rapportversion" in text
    assert "Faglig reviewer" in text and saved["editorial_revision"]["note"] in text
    assert (
        "Tidligere JEV-kontrol — teksten er ændret og kræver nyt kvalitetstjek" in text
    )
    assert saved["editorial_revision"]["changes"][0]["before"]["text"] in text
    with ZipFile(BytesIO(export_dpia_xlsx(setup["request"], result))) as book:
        xml = book.read("xl/worksheets/ai-provenance.xml").decode()
        assert "Fagligt redigeret rapportversion" in xml
        assert "Faglig reviewer" in xml and "Tidligere JEV-kontrol" in xml


def test_human_only_report_gets_export_provenance_without_fake_ai_claim(setup):
    with setup["factory"]() as db:
        record = db.get(DPIAAssessmentRecord, setup["base"]["id"])
        record.result_payload = {**record.result_payload, "ai_generation": None}
        db.commit()
    body = payload(setup)
    body["changes"][0]["source_ids"] = []
    response = setup["client"].post(endpoint(setup), json=body)
    assert response.status_code == 201, response.text
    result = DPIAAssessmentResponse.model_validate(response.json())
    assert result.ai_generation is None
    doc = Document(BytesIO(export_dpia_docx(setup["request"], result)))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Den redigerede tekst er ikke kontrolleret med JEV." in text
    assert "tidligere JEV-kontrol" not in text
    with ZipFile(BytesIO(export_dpia_xlsx(setup["request"], result))) as book:
        xml = book.read("xl/worksheets/ai-provenance.xml").decode()
        assert "Faglig reviewer" in xml
        assert "Den redigerede tekst er ikke kontrolleret med JEV." in xml
        assert "tidligere JEV-kontrol" not in xml


def test_summary_revision_export_uses_danish_field_names(setup):
    body = payload(setup, changes=[{
        "kind": "summary", "target_id": "summary", "field": "executive_summary",
        "text": "Resuméet er fagligt præciseret med kommunens konkrete afgrænsning.",
        "source_ids": [],
    }])
    saved = setup["client"].post(endpoint(setup), json=body)
    assert saved.status_code == 201, saved.text
    result = DPIAAssessmentResponse.model_validate(saved.json())
    doc = Document(BytesIO(export_dpia_docx(setup["request"], result)))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Ændrede afsnit: Resumé" in text and "Ændring: Resumé" in text
    assert "executive_summary" not in text and "Ændrede afsnit: summary" not in text
    with ZipFile(BytesIO(export_dpia_xlsx(setup["request"], result))) as book:
        xml = book.read("xl/worksheets/ai-provenance.xml").decode()
        assert "Resumé" in xml and "executive_summary" not in xml


def test_unknown_assessment_does_not_reveal_other_history(setup):
    url = endpoint(setup, str(uuid4()))
    assert setup["client"].get(url).status_code == 404
    assert setup["client"].post(url, json=payload(setup)).status_code == 404


def test_noop_and_conflicting_shared_sources_cannot_create_versions(setup):
    original = next(
        item for item in setup["base"]["sections"] if item["id"] == setup["section"]
    )
    body = payload(setup)
    body["changes"][0]["text"] = original["text"]
    assert setup["client"].post(endpoint(setup), json=body).status_code == 422
    body = payload(
        setup,
        changes=[
            {
                "kind": "summary",
                "target_id": "summary",
                "field": "executive_summary",
                "text": "Nyt resumé",
                "source_ids": [],
            },
            {
                "kind": "scope",
                "target_id": "scope",
                "field": "scope",
                "text": "Ny afgrænsning",
                "source_ids": ["input:purpose"],
            },
        ],
    )
    assert setup["client"].post(endpoint(setup), json=body).status_code == 422


def test_routes_require_workspace_access(setup):
    from fastapi import HTTPException

    def denied():
        raise HTTPException(403, "Ingen adgang")

    setup["app"].dependency_overrides[api.WORKSPACE_ACCESS] = denied
    assert setup["client"].get(endpoint(setup)).status_code == 403
    assert setup["client"].post(endpoint(setup), json=payload(setup)).status_code == 403
