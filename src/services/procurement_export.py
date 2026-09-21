"""Readable Word brief of an immutable procurement fact-review snapshot."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
import re
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from src.services.ai_presentation import metadata_identifier, metadata_note, model_label
from src.services.dpia_assessment import (
    CONTROL_LABELS,
    DATA_LABELS,
    HOSTING_LABELS,
    SOLUTION_LABELS,
    SUBJECT_LABELS,
)


FIELD_LABELS = {
    "purpose": "Formålsbeskrivelse i materialet",
    "processing_description": "Behandling af oplysninger",
    "supplier_name": "Leverandør",
    "solution_type": "Løsningstype",
    "hosting_region": "Placering af data",
    "transfer_outside_eea": "Overførsel uden for EU og EØS",
    "model_training": "Brug af oplysninger til AI træning",
    "retention_period": "Opbevaringsperiode",
    "data_subjects": "Grupper af registrerede",
    "personal_data_categories": "Kategorier af personoplysninger",
    "special_categories": "Følsomme personoplysninger",
    "criminal_data": "Oplysninger om strafbare forhold",
    "cpr_data": "CPR numre",
    "vulnerable_subjects": "Børn eller andre sårbare registrerede",
    "large_scale": "Behandling i stort omfang",
    "systematic_monitoring": "Systematisk overvågning",
    "automated_decisions": "Automatiske afgørelser",
    "human_oversight": "Menneskeligt tilsyn",
    "controls": "Oplyste eller foreslåede foranstaltninger",
    "secondary_uses": "Andre anvendelser",
}
VALUE_LABELS = {
    "data_subjects": SUBJECT_LABELS,
    "personal_data_categories": DATA_LABELS,
    "solution_type": SOLUTION_LABELS,
    "hosting_region": HOSTING_LABELS,
    "controls": CONTROL_LABELS,
}


def clean_text(value: object) -> str:
    return "".join(
        char
        for char in ("" if value is None else str(value))
        if char in "\t\n\r"
        or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD
        or 0x10000 <= ord(char) <= 0x10FFFF
    )


def fact_value(field: str, value: Any) -> str:
    if value is None:
        return "Ikke afklaret"
    if isinstance(value, bool):
        return "Ja" if value else "Nej"
    labels = VALUE_LABELS.get(field, {})
    if isinstance(value, list):
        return (
            "; ".join(labels.get(item, str(item)) for item in value) or "Ingen angivet"
        )
    return labels.get(value, str(value))


def date_label(value: str | None) -> str:
    if not value:
        return "Ikke registreret"
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        parsed = parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)
        return parsed.astimezone(ZoneInfo("Europe/Copenhagen")).strftime(
            "%d.%m.%Y kl. %H.%M"
        )
    except ValueError:
        return clean_text(value)


def paragraph(doc, value: object, *, label: str | None = None, style=None):
    item = doc.add_paragraph(style=style)
    if label:
        item.add_run(clean_text(label) + ": ").bold = True
    item.add_run(clean_text(value))
    return item


def references(doc, refs: list[dict], sources: dict[str, dict]):
    for ref in refs:
        source = sources.get(ref["source_id"], {})
        locator = source.get("locator") or "Placering ikke angivet"
        quote = paragraph(doc, ref["quote"], style="Quote")
        quote.paragraph_format.space_after = Pt(3)
        paragraph(
            doc,
            f"{source.get('title', ref['source_id'])} · {locator} · version {source.get('version', 'ukendt')}",
            label="Kilde",
        )
        if source.get("source_url"):
            paragraph(doc, source["source_url"], label="Hjemmeside")


def keep_short_block(doc, start: int):
    items = doc.paragraphs[start:]
    if sum(len(item.text) for item in items) < 2500:
        for item in items[:-1]:
            item.paragraph_format.keep_with_next = True


def number_ranges(values: list[int]) -> str:
    ordered = sorted(set(values))
    result = []
    for number in ordered:
        if result and result[-1][1] == number - 1:
            result[-1][1] = number
        else:
            result.append([number, number])
    return ", ".join(
        str(start) if start == end else f"{start}–{end}" for start, end in result
    )


def grouped_sources(sources: list[dict]) -> list[dict]:
    """One source record per stored document version, preserving every locator."""
    groups = {}
    for source in sources:
        key = tuple(
            source.get(field)
            for field in (
                "document_version_id",
                "version",
                "checksum",
                "title",
                "source_url",
            )
        )
        if key not in groups:
            groups[key] = {"source": source, "locators": [], "ids": []}
        groups[key]["locators"].append(source.get("locator", "Placering ikke angivet"))
        groups[key]["ids"].append(source["id"])
    result = []
    for group in groups.values():
        locators, numbered = [], {}
        for locator in dict.fromkeys(group["locators"]):
            match = re.fullmatch(r"(.+?) (\d+)", locator)
            if match:
                numbered.setdefault(match[1], []).append(int(match[2]))
            else:
                locators.append(locator)
        group["locator_label"] = "; ".join(
            [
                *(
                    f"{prefix} {number_ranges(numbers)}"
                    for prefix, numbers in numbered.items()
                ),
                *locators,
            ]
        )
        source = group["source"]
        prefix = f"document:{source.get('document_version_id')}:"
        if all(
            item.startswith(prefix) and item[len(prefix) :].isdigit()
            for item in group["ids"]
        ):
            group["reference_label"] = (
                f"{prefix} · uddrag {number_ranges([int(item[len(prefix):]) for item in group['ids']])}"
            )
        else:
            group["reference_label"] = "; ".join(group["ids"])
        result.append(group)
    return result


def build_procurement_review_docx(
    analysis: dict, review: dict, *, outdated: bool = False, case_reference: str = ""
) -> bytes:
    """Use only persisted snapshots; never fetch websites or call AI here."""
    if (
        analysis["id"] != review["analysis_id"]
        or analysis["case_id"] != review["case_id"]
    ):
        raise ValueError("Gennemgangen og analysen tilhører ikke samme sag.")
    profile = analysis["profile_snapshot"]
    prefill = review["dpia_prefill"]
    sources = {source["id"]: source for source in analysis["sources"]}
    facts = {fact["id"]: fact for fact in analysis["facts"]}
    accepted = review["accepted_fact_ids"]
    if len(set(accepted)) != len(accepted) or any(
        fact_id not in facts for fact_id in accepted
    ):
        raise ValueError("Den gemte gennemgang indeholder ukendte oplysninger.")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(21), Mm(20)
    section.left_margin, section.right_margin = Mm(22), Mm(22)
    for name in (
        "Normal",
        "Title",
        "Subtitle",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Quote",
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.widow_control = True
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(11.5)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Quote"].font.size = Pt(10)
    doc.styles["Quote"].paragraph_format.left_indent = Mm(5)
    doc.styles["Quote"].paragraph_format.right_indent = Mm(3)
    doc.core_properties.title = (
        f"Grundlag for juridisk gennemgang af {profile['system_name']}"
    )
    doc.core_properties.subject = (
        "Vurderingsgrundlag for AI-løsninger og dokumenterede afklaringspunkter"
    )
    doc.core_properties.author = "S.H.I.E.L.D."
    doc.core_properties.comments = ""
    doc.add_paragraph("Grundlag for juridisk gennemgang", style="Title")
    doc.add_paragraph(clean_text(profile["system_name"]), style="Subtitle")
    paragraph(
        doc,
        "Materialet samler kommunens påtænkte anvendelse, de oplysninger som sagsbehandleren har valgt til vurderingsgrundlaget, og de spørgsmål som kræver faglig eller juridisk afklaring. Brug det til dialog mellem systemejer, jura, informationssikkerhed og databeskyttelsesrådgiver.",
    )
    paragraph(
        doc,
        "Vurderingsarbejdet omfatter AI-løsninger og IT-løsninger med AI-funktioner. Kilderne skal beskrive den konkrete AI-funktion, dens opgave, input, output og kommunens menneskelige kontrol. Hvis AI-funktionen ikke er dokumenteret, er den uafklaret; dette dokument bekræfter ikke i sig selv, at løsningen indeholder AI.",
        label="AI-afgrænsning",
    )
    paragraph(
        doc,
        "Kræver faglig og juridisk gennemgang. Oplysningerne er gennemgået som vurderingsgrundlag; anskaffelsen, behandlingsgrundlaget og risiciene er ikke juridisk godkendt med dette dokument.",
        label="Status",
    )
    if outdated:
        paragraph(
            doc,
            "Sagens oplysninger eller kilder er ændret siden denne analyse. Dokumentet gengiver den historiske version, der blev gennemgået. Gennemgå det aktuelle grundlag, før oplysningerne anvendes videre.",
            label="Historisk grundlag",
        )
    doc.add_heading("Kommunens anvendelse og vurderingsbehov", level=1)
    stages = {
        "new_purchase": "Ny anskaffelse",
        "renewal": "Kontraktfornyelse",
        "change": "Ændret anvendelse",
    }
    for label, value in [
        ("Kommune", prefill.get("organisation", profile.get("organisation"))),
        ("Fagområde", prefill.get("department", profile.get("department"))),
        ("Systemejer", prefill.get("owner", profile.get("owner"))),
        (
            "Leverandør",
            prefill.get("supplier_name", profile.get("supplier_name")) or "Ikke oplyst",
        ),
        ("Anledning", stages.get(profile.get("procurement_stage"), "Ikke oplyst")),
        (
            "Journalreference",
            profile.get("journal_reference") or case_reference or review["case_id"],
        ),
    ]:
        paragraph(doc, value, label=label)
    paragraph(
        doc,
        prefill.get("purpose", profile["intended_use"]),
        label="Påtænkt kommunal anvendelse",
    )
    doc.add_heading("Resumé af kildematerialet", level=1)
    paragraph(doc, analysis["summary"])
    paragraph(
        doc,
        "Leverandørmateriale dokumenterer leverandørens udsagn. Det bekræfter ikke i sig selv kommunens konkrete anvendelse, aftalens indgåelse eller den faktiske drift.",
    )
    doc.add_heading("Oplysninger valgt til vurderingsgrundlaget", level=1)
    paragraph(
        doc,
        f"{len(accepted)} af {len(facts)} udtrukne oplysninger er valgt af sagsbehandleren. De øvrige oplysninger er ikke medtaget i forudfyldningen af konsekvensanalysen.",
    )
    for index, fact_id in enumerate(accepted, 1):
        block_start = len(doc.paragraphs)
        fact = facts[fact_id]
        doc.add_heading(
            f"{index} {FIELD_LABELS.get(fact['field'], fact.get('label') or fact['field'])}",
            level=2,
        )
        paragraph(doc, fact_value(fact["field"], fact["value"]), label="Oplysning")
        if fact["field"] == "purpose":
            paragraph(
                doc,
                "Formålsbeskrivelsen i materialet er baggrundsviden. Kommunens påtænkte anvendelse ovenfor er anvendt i vurderingsgrundlaget.",
            )
        if fact["field"] == "controls":
            paragraph(
                doc,
                "Oplyst eller foreslået foranstaltning. Den er ikke registreret som verificeret implementeret kontrol.",
            )
        references(doc, fact["source_refs"], sources)
        keep_short_block(doc, block_start)
    if not accepted:
        paragraph(
            doc,
            "Ingen udtrukne oplysninger er valgt. Vurderingen er kun forberedt med kommunens egne profiloplysninger.",
        )
    doc.add_heading("JEV markeringer til gennemgang", level=1)
    checks = analysis["review"]["checks"]
    flagged = [check for check in checks if check["requires_review"]]
    paragraph(
        doc,
        f"JEV har kontrolleret {len(checks)} punkter og markeret {len(flagged)} til nærmere gennemgang. Kontrollen er en støtte til kildegennemgang og er ikke juridisk godkendelse eller en måling af juridisk sikkerhed.",
    )
    for index, check in enumerate(flagged, 1):
        item_id = check["id"].removeprefix("fact:")
        fact = facts.get(item_id) if check["id"].startswith("fact:") else None
        label = (
            FIELD_LABELS.get(fact["field"], check["label"])
            if fact
            else (
                "Resumé af kildematerialet"
                if check["id"] == "summary"
                else check["label"]
            )
        )
        doc.add_heading(f"{index} {clean_text(label)}", level=2)
        paragraph(doc, check["id"], label="Kontrolreference")
        if fact:
            paragraph(
                doc,
                fact_value(fact["field"], fact["value"]),
                label=FIELD_LABELS.get(fact["field"], fact["field"]),
            )
            paragraph(
                doc,
                (
                    "Valgt af sagsbehandleren med fagligt notat"
                    if item_id in accepted
                    else "Ikke valgt til vurderingsgrundlaget"
                ),
                label="Sagsbehandlerens valg",
            )
            references(doc, fact["source_refs"], sources)
        elif check["id"] == "summary":
            paragraph(doc, analysis["summary"])
    if not flagged:
        paragraph(
            doc,
            "Ingen punkter er markeret af JEV. Den faglige og juridiske gennemgang er fortsat nødvendig.",
        )
    doc.add_heading("Modstridende oplysninger", level=1)
    for index, conflict in enumerate(analysis["conflicts"], 1):
        doc.add_heading(f"Afklaringspunkt {index}", level=2)
        paragraph(doc, conflict["description"])
        references(doc, conflict["source_refs"], sources)
    if not analysis["conflicts"]:
        paragraph(
            doc,
            "Analysen har ikke udpeget modstridende oplysninger. Det dokumenterer ikke, at materialet er fuldstændigt.",
        )
    doc.add_heading("Spørgsmål til faglig og juridisk afklaring", level=1)
    for index, question in enumerate(analysis["questions"], 1):
        priority = (
            "Høj prioritet" if question["priority"] == "high" else "Normal prioritet"
        )
        doc.add_heading(f"{index} {clean_text(question['topic'])}", level=2)
        paragraph(doc, question["question"])
        paragraph(doc, priority, label="Prioritet")
    doc.add_heading("Sagsbehandlerens notat", level=1)
    paragraph(doc, review.get("note") or "Der er ikke tilføjet et notat.")
    paragraph(doc, date_label(review["created_at"]), label="Gennemgået")
    paragraph(doc, review["reviewed_by"], label="Registreret brugerreference")
    doc.add_heading("Kildeoversigt og versionsspor", level=1)
    for index, group in enumerate(grouped_sources(analysis["sources"]), 1):
        block_start = len(doc.paragraphs)
        source = group["source"]
        doc.add_heading(f"{index} {clean_text(source['title'])}", level=2)
        paragraph(
            doc,
            f"{group['locator_label']} · version {source.get('version', 'ukendt')}",
        )
        if source.get("source_url"):
            paragraph(doc, source["source_url"], label="Hjemmeside")
        paragraph(doc, group["reference_label"], label="Kildereference")
        paragraph(
            doc,
            source.get("document_version_id", "Ikke registreret"),
            label="Dokumentversion",
        )
        paragraph(doc, source.get("checksum", "Ikke registreret"), label="Kontrolsum")
        keep_short_block(doc, block_start)
    doc.add_heading("Analyse og gennemgang", level=1)
    for label, value in [
        ("Sag", review["case_id"]),
        ("Analyse", analysis["id"]),
        ("Analysetidspunkt", date_label(analysis["created_at"])),
        ("Gennemgang", review["id"]),
        ("Skrivemodel", model_label(analysis["model"])),
        ("Kontrolmodel", analysis["review"]["model"]),
        (
            "Versionsgrundlag",
            " · ".join(
                filter(
                    None,
                    [
                        metadata_identifier(analysis.get("prompt_version")),
                        analysis["review"]["rubric_version"],
                    ],
                )
            ),
        ),
    ]:
        item = paragraph(doc, value, label=label)
        item.paragraph_format.space_after = Pt(3)
        for run in item.runs:
            run.font.size = Pt(9)
    if analysis.get("provenance_note"):
        paragraph(doc, metadata_note(analysis["provenance_note"]))
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(4)
    footer.add_run("Grundlag for juridisk gennemgang · Side ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
