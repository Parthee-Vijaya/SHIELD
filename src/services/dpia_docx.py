"""Readable Danish Word export of the persisted DPIA assessment snapshot.

The document follows the supplied consequence-analysis structure without
bundling its organisation-specific facts. Export does not call a model, fetch
sources, accept risks, or infer an approval from an assessment status.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from io import BytesIO
import re
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .dpia_assessment import (
    CONTROL_LABELS,
    DPIAAssessmentRequest,
    DPIAAssessmentResponse,
    risk_matrix,
)
from .dpia_export import (
    APPROVAL_TYPE_LABELS,
    IMPACT_LABELS,
    LIKELIHOOD_LABELS,
    RISK_LEVEL_LABELS,
    RISK_ROWS,
    SECTION_ROWS,
    _validate_snapshot_contract,
    _editorial_change_label,
    _editorial_review_note,
    _editorial_target_label,
)
from .dpia_reading_guide import reading_guide

RISK_COLORS = {
    "low": "C6EFCE",
    "medium": "FFEB9C",
    "high": "FFC000",
    "very_high": "C00000",
}
REVIEW_LABELS = {
    "requires_review": "Kræver faglig gennemgang",
    "missing_information": "Mangler oplysninger",
    "not_applicable": "Ikke relevant ifølge vurderingen",
    "requires_verification": "Kræver verificering",
    "findings_require_review": "Fund kræver faglig gennemgang",
    "requires_human_review": "Kræver menneskelig gennemgang",
}
SOURCE_LABELS = {
    "provided_input": "Oplysninger fra sagen",
    "deterministic_rule": "Regelbaseret vurdering",
    "missing_information": "Manglende oplysninger",
    "ai_generated": "AI-genereret udkast",
    "ai_assisted": "AI-assisteret udkast",
}
RISK_GROUPS = {
    "3": "Afgrænsning af aktivitet og anvendelse",
    "4": "Udvikling",
    "5": "Test",
    "6": "Drift og monitorering",
}
LIKELIHOOD_DESCRIPTIONS = {
    1: "Hændelsen anses for næsten udelukket og kendes kun fra få uafhængige tilfælde.",
    2: "Hændelsen forventes ikke at forekomme, men kendes fra andre organisationer.",
    3: "Hændelsen er sandsynlig og kendes fra erfaring eller tilbagevendende hændelser.",
    4: "Hændelsen forventes at forekomme og opleves jævnligt i sammenlignelige organisationer.",
}
IMPACT_DESCRIPTIONS = {
    1: "Få uhensigtsmæssigheder, der kan overkommes uden større indsats.",
    2: "Betydelige uhensigtsmæssigheder, der kan overkommes med en indsats.",
    3: "Betydelige følger, som kun kan overkommes med væsentlig indsats og konsekvenser for den enkelte.",
    4: "Indgribende følger, som kun vanskeligt eller slet ikke kan overkommes.",
}


def _text(value: object) -> str:
    """Keep user/model content as plain text and remove invalid XML controls."""
    text = "" if value is None else str(value)
    return "".join(
        c
        for c in text
        if c in "\t\n\r"
        or 0x20 <= ord(c) <= 0xD7FF
        or 0xE000 <= ord(c) <= 0xFFFD
        or 0x10000 <= ord(c) <= 0x10FFFF
    )


def _paragraph(
    doc: DocumentType, text: object, *, label: str = "", keep_with_next: bool = False
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep_with_next
    if label:
        p.add_run(label + (" " if label.endswith(("?", ":")) else ": ")).bold = True
    p.add_run(_text(text))


def _structured_text(
    doc: DocumentType, text: object, *, label: str = "", keep_with_next: bool = False
) -> None:
    """Render authored paragraphs and lists without interpreting their claims."""
    if label:
        _paragraph(doc, "", label=label, keep_with_next=True)
    for block in re.split(r"\n\s*\n", _text(text).strip()):
        lines = block.splitlines()
        if lines and all(re.match(r"^\s*(?:[-*•]|\d+[.)])\s+", line) for line in lines):
            for line in lines:
                numbered = bool(re.match(r"^\s*\d+[.)]\s+", line))
                doc.add_paragraph(
                    re.sub(r"^\s*(?:[-*•]|\d+[.)])\s+", "", line),
                    style="List Number" if numbered else "List Bullet",
                )
        else:
            _paragraph(doc, block)
    if keep_with_next:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True


def _bullet_list(doc: DocumentType, values: Iterable[object]) -> None:
    for value in values:
        doc.add_paragraph(_text(value), style="List Bullet")


def _recommendations(doc: DocumentType, guide: dict) -> None:
    recommendations = guide.get("recommendations", [])
    _heading(doc, "Anbefalinger – forslag til faglig drøftelse", 2)
    _paragraph(
        doc,
        "Forslagene er adskilt fra vurderingen. De dokumenterer ikke gennemførte "
        "foranstaltninger, godkendelse eller lavere risiko. Eventuelle ændringer "
        "skal besluttes, gennemføres og verificeres, før risikoen vurderes på ny.",
    )
    _paragraph(
        doc,
        (
            "Gemte forslag i denne vurderingsversion."
            if guide.get("recommendation_origin") == "saved"
            else "Regelbaserede forslag ud fra sagens oplysninger. De er en læsehjælp og indgår ikke i det historiske AI-udkast eller dets JEV-kontrol."
        ),
        label="Anbefalingernes grundlag",
    )
    if not recommendations:
        _paragraph(
            doc,
            "Der er ikke registreret særskilte anbefalinger i denne vurderingsversion.",
        )
    for item in recommendations:
        _heading(doc, _text(item.get("title") or "Anbefaling"), 3)
        for key, label in (
            ("proposal", "Anbefalet mulighed"),
            ("rationale", "Hvorfor kan forslaget være relevant?"),
            ("prerequisites", "Forudsætninger før et valg"),
            ("verification", "Sådan efterprøves forslaget"),
        ):
            value = item.get(key)
            if value:
                if isinstance(value, list):
                    _paragraph(doc, "", label=label, keep_with_next=True)
                    _bullet_list(doc, value)
                else:
                    _structured_text(doc, value, label=label)
        _source_ids(doc, item.get("source_ids", []))


def _heading(
    doc: DocumentType, title: str, level: int = 1, *, new_page: bool = False
) -> None:
    p = doc.add_heading(_text(title), level=level)
    p.paragraph_format.page_break_before = new_page


def _fill(cell: Any, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _table(
    doc: DocumentType,
    headers: list[str],
    rows: Iterable[Iterable[object]],
    widths: list[float],
    *,
    keep_together: bool = False,
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = _text(value)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{attr}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                child = OxmlElement(f"w:{side}")
                child.set(qn("w:w"), "80")
                child.set(qn("w:type"), "dxa")
                margins.append(child)
            cell._tc.get_or_add_tcPr().append(margins)
            _fill(
                cell,
                "1F3864" if index == 0 else ("F4F6F8" if index % 2 == 0 else "FFFFFF"),
            )
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.keep_with_next = index == 0 or (
                    keep_together and index < len(table.rows) - 1
                )
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _risk_fill(cell: Any, level: str) -> None:
    _fill(cell, RISK_COLORS[level])
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(
            "FFFFFF" if level == "very_high" else "000000"
        )


def _source_ids(doc: DocumentType, values: object) -> None:
    if isinstance(values, (list, tuple)) and values:
        _paragraph(
            doc, ", ".join(_text(value) for value in values), label="Kildehenvisninger"
        )


def _document(request: DPIAAssessmentRequest) -> DocumentType:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = (
        section.right_margin
    ) = Mm(20)
    section.header_distance = section.footer_distance = Mm(9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in (
        ("Title", 23, "000000"),
        ("Heading 1", 15, "1F3864"),
        ("Heading 2", 12, "2E5496"),
        ("Heading 3", 11, "000000"),
    ):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    # The default Word template can carry a decorative title border.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_name = _text(request.project_name)
    if len(header_name) > 96:
        header_name = header_name[:93] + "..."
    run = header.add_run(f"Konsekvensanalyse · {header_name}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Side ")
    for field_name, separator in (("PAGE", " af "), ("NUMPAGES", "")):
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        footer._p.append(field)
        footer.add_run(separator)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    return doc


def _criteria(doc: DocumentType) -> None:
    _heading(doc, "4 Evalueringskriterier", new_page=True)
    _heading(doc, "Sandsynlighed", 2)
    _table(
        doc,
        ["Niveau", "Beskrivelse"],
        [(LIKELIHOOD_LABELS[i], LIKELIHOOD_DESCRIPTIONS[i]) for i in range(1, 5)],
        [50, 120],
    )
    _heading(doc, "Konsekvens for de registrerede", 2)
    _table(
        doc,
        ["Niveau", "Beskrivelse"],
        [(IMPACT_LABELS[i], IMPACT_DESCRIPTIONS[i]) for i in range(1, 5)],
        [50, 120],
    )
    _heading(doc, "Risikomatrix", 2)
    _paragraph(
        doc,
        "Risikoniveauet findes i skabelonens matrix ud fra sandsynlighed og konsekvens. "
        "Samme matrix anvendes før og efter foranstaltninger.",
    )
    table = _table(
        doc,
        ["Konsekvens / sandsynlighed"] + [str(i) for i in range(1, 5)],
        [
            [IMPACT_LABELS[impact]]
            + [
                RISK_LEVEL_LABELS[risk_matrix(likelihood, impact)]
                for likelihood in range(1, 5)
            ]
            for impact in range(4, 0, -1)
        ],
        [58, 28, 28, 28, 28],
    )
    for row, impact in zip(table.rows[1:], range(4, 0, -1)):
        for cell, likelihood in zip(row.cells[1:], range(1, 5)):
            _risk_fill(cell, risk_matrix(likelihood, impact))


def _risk_register(doc: DocumentType, result: DPIAAssessmentResponse) -> None:
    _heading(doc, "5 Risici", new_page=True)
    _paragraph(
        doc,
        "Hver risiko viser den iboende risiko og den resterende risiko med de "
        "foranstaltninger, der er lagt til grund i denne vurderingsversion. Planlagte "
        "foranstaltninger er ikke i sig selv dokumentation for en lavere risiko.",
    )
    indexed = {risk.id: risk for risk in result.risks}
    group = None
    for risk_id in RISK_ROWS:
        risk = indexed[risk_id]
        if risk_id.split(".")[0] != group:
            group = risk_id.split(".")[0]
            _heading(doc, RISK_GROUPS[group], 2)
        _heading(doc, f"{risk.id} {risk.area}", 3)
        _structured_text(doc, risk.scenario, label="Hvad kan ske?")
        _structured_text(
            doc,
            getattr(risk, "rationale", "")
            or "En særskilt begrundelse er ikke dokumenteret i denne vurderingsversion og skal afklares fagligt.",
            label="Hvorfor er dette en risiko?",
        )
        _structured_text(
            doc,
            getattr(risk, "consequences", "")
            or "De konkrete følger for de registrerede er ikke særskilt beskrevet og skal afklares fagligt.",
            label="Hvem rammes – og hvad er konsekvensen?",
        )
        table = _table(
            doc,
            ["Vurdering", "Sandsynlighed", "Konsekvens", "Risiko"],
            [
                [
                    "Iboende før foranstaltninger",
                    LIKELIHOOD_LABELS[risk.likelihood],
                    IMPACT_LABELS[risk.impact],
                    RISK_LEVEL_LABELS[risk.inherent_risk],
                ],
                [
                    "Efter foranstaltninger",
                    LIKELIHOOD_LABELS[risk.residual_likelihood],
                    IMPACT_LABELS[risk.residual_impact],
                    RISK_LEVEL_LABELS[risk.residual_risk],
                ],
            ],
            [47, 43, 43, 37],
            keep_together=True,
        )
        _risk_fill(table.rows[1].cells[3], risk.inherent_risk)
        _risk_fill(table.rows[2].cells[3], risk.residual_risk)
        _structured_text(
            doc, risk.measures, label="Mitigerende forslag og oplyste foranstaltninger"
        )
        _paragraph(
            doc,
            "Et forslag er ikke dokumentation for gennemførelse. Kontrollér effekten "
            "og dokumentér evidensen før en eventuel ændring af den resterende risiko.",
            label="Hvad skal verificeres?",
        )
        status = REVIEW_LABELS.get(
            risk.implementation_status, risk.implementation_status
        )
        _paragraph(
            doc,
            f"{risk.owner or 'Ikke angivet'} · Status: {status} · Frist: {risk.due_date or 'Ikke fastsat'}",
            label="Risikoejer",
            keep_with_next=bool(getattr(risk, "source_ids", [])),
        )
        _source_ids(doc, getattr(risk, "source_ids", []))
    ai = getattr(result, "ai_generation", None) or {}
    additional = getattr(result, "additional_risks", None) or ai.get(
        "additional_risks", []
    )
    if additional:
        _heading(doc, "Yderligere foreslåede risici", 2)
        _paragraph(
            doc,
            "Disse forslag indgår ikke i de 33 skabelonrisici eller den beregnede "
            "samlede risiko. De skal vurderes og tildeles ansvar og risikoniveau.",
        )
        for index, risk in enumerate(additional, 1):
            if not isinstance(risk, Mapping):
                continue
            _heading(
                doc, f"F{index} {_text(risk.get('title') or 'Yderligere risiko')}", 3
            )
            for key, label in (
                ("scenario", "Hvad kan ske?"),
                ("measures", "Foreslåede foranstaltninger"),
                ("rationale", "Hvorfor er dette en risiko?"),
            ):
                if risk.get(key):
                    _structured_text(doc, risk[key], label=label)
            _source_ids(doc, risk.get("source_ids", []))


def _stakeholders(
    doc: DocumentType,
    request: DPIAAssessmentRequest,
    result: DPIAAssessmentResponse,
    guide: dict,
) -> None:
    _heading(doc, "7 Interessentinddragelse og godkendelse", new_page=True)
    _heading(doc, "Databeskyttelsesrådgiverens synspunkter", 2)
    _paragraph(
        doc,
        ("DPO/databeskyttelsesrådgiverens inddragelse er ikke dokumenteret og skal afklares. " + request.dpo_advice)
        if request.dpo_involved is None else request.dpo_advice
        or (
            "DPO er oplyst som inddraget, men konkrete synspunkter er ikke dokumenteret i vurderingen."
            if request.dpo_involved
            else "DPO er ikke oplyst som inddraget. Rådgivning skal dokumenteres."
        ),
    )
    _heading(doc, "Registrerede eller deres repræsentanter", 2)
    _paragraph(
        doc,
        request.data_subject_consultation
        or "Synspunkter eller en begrundelse for fravalg af inddragelse er ikke dokumenteret.",
    )
    _heading(doc, "Høring af Datatilsynet", 2)
    _paragraph(
        doc,
        "Behovet for forudgående høring efter artikel 36 skal vurderes på baggrund "
        "af den resterende risiko og dokumenterede foranstaltninger. Denne eksport "
        "dokumenterer ikke en beslutning om at gennemføre eller undlade høring.",
    )
    _heading(doc, "Ledelsesgodkendelse", 2)
    approval = guide.get("approval", {})
    if approval.get("items"):
        _paragraph(doc, approval.get("label", "Se den registrerede beslutning."))
        _paragraph(
            doc,
            "Beslutningens type og vilkår afgør, hvad der er taget stilling til. Kontrollér, om den omfatter ledelsesgodkendelse og accept af resterende risiko.",
        )
        for item in approval["items"]:
            _paragraph(
                doc,
                APPROVAL_TYPE_LABELS.get(item.get("approval_type"), "Ikke angivet"),
                label="Beslutningstype",
            )
            _paragraph(
                doc, item.get("decided_by") or "Ikke angivet", label="Besluttet af"
            )
            _paragraph(
                doc, item.get("decided_at") or "Ikke angivet", label="Beslutningsdato"
            )
            for key, label in (("reason", "Begrundelse"), ("conditions", "Vilkår")):
                if item.get(key):
                    value = item[key]
                    _structured_text(
                        doc,
                        "\n".join(value) if isinstance(value, list) else value,
                        label=label,
                    )
            if item.get("is_identity_verified") is not True:
                _paragraph(
                    doc,
                    "Beslutningens identitet er ikke verificeret og skal bekræftes før den lægges til grund.",
                )
    else:
        _paragraph(
            doc,
            "Formel ledelsesgodkendelse og accept af resterende risiko er ikke dokumenteret "
            "i denne vurderingsversion. Vurderingsstatus er ikke en godkendelse.",
        )
    _heading(doc, "Ajourføring og offentliggørelse", 2)
    _paragraph(doc, request.owner, label="Oplyst ansvarlig")
    _paragraph(
        doc,
        request.publication_plan or "Plan for offentliggørelse er ikke dokumenteret.",
    )
    for title, values in (
        ("Blokerende forhold", result.blockers),
        ("Manglende oplysninger", result.missing_information),
        ("Næste skridt", result.next_steps),
    ):
        _heading(doc, title, 2)
        _bullet_list(doc, values or ["Ingen registreret i denne vurderingsversion."])
    ai = getattr(result, "ai_generation", None) or {}
    questions = list(getattr(result, "open_questions", []) or []) + list(
        ai.get("open_questions", []) or []
    )
    if questions:
        _heading(doc, "Åbne spørgsmål", 2)
        _bullet_list(doc, dict.fromkeys(_text(value) for value in questions))


def _provenance(doc: DocumentType, result: DPIAAssessmentResponse) -> None:
    _heading(doc, "Kilder og vurderingsgrundlag", 1, new_page=True)
    ai = getattr(result, "ai_generation", None) or {}
    editorial = result.editorial_revision or {}
    stale_checks = set(editorial.get("stale_check_ids", []))
    if editorial:
        _heading(doc, "Fagligt redigeret rapportversion", 2)
        _paragraph(
            doc, editorial.get("edited_by", "Ikke registreret"), label="Redigeret af"
        )
        _paragraph(doc, editorial.get("edited_at", ""), label="Redigeret")
        _paragraph(doc, editorial.get("note", ""), label="Begrundelse")
        _paragraph(
            doc, editorial.get("base_assessment_id", ""), label="Forrige version"
        )
        _paragraph(
            doc,
            ", ".join(
                _editorial_target_label(target)
                for target in editorial.get("changed_targets", [])
            ),
            label="Ændrede afsnit",
        )
        _paragraph(
            doc,
            "Teksten er fagligt redigeret. "
            + _editorial_review_note(result)
            + " Risikoscorer, blokeringer, sagsoplysninger og hjemmelskontrol er uændrede; versionen er ikke en juridisk godkendelse.",
        )
        for change in editorial.get("changes", []):
            _paragraph(doc, _editorial_change_label(change), label="Ændring")
            _paragraph(
                doc, change.get("before", {}).get("text", ""), label="Tidligere tekst"
            )
            _paragraph(doc, change.get("after", {}).get("text", ""), label="Ny tekst")
    if ai:
        _paragraph(doc, ai.get("model", "Ikke registreret"), label="Anvendt model")
        if ai.get("provider") in {"codex-local-test", "codex-local"}:
            is_test = ai.get("provider") == "codex-local-test"
            _paragraph(
                doc,
                "Codex – lokal testkørsel" if is_test else "Codex – lokal kørsel",
                label="Udarbejdet i",
            )
            if ai.get("run_id"):
                _paragraph(
                    doc, ai["run_id"], label="Testkørsel" if is_test else "Kørsels-ID"
                )
            if ai.get("model_run_provenance") == "operator_reported":
                _paragraph(
                    doc,
                    "Model og kørsels-ID er oplyst af den lokale operatør ved importen.",
                    label="Oplysning om model og kørsel",
                )
        if ai.get("prompt_version"):
            _paragraph(doc, ai["prompt_version"], label="Instruktionsversion")
        review = ai.get("review") or ai.get("evaluation")
        if isinstance(review, Mapping):
            for key, label in (
                ("status", "Kontrolstatus"),
                ("summary", "Kontrolresultat"),
                ("conclusion", "Kontrolkonklusion"),
            ):
                if review.get(key):
                    _paragraph(
                        doc,
                        REVIEW_LABELS.get(str(review[key]), review[key]),
                        label=label,
                    )
            for key, label in (
                ("model", "Kontrolmodel"),
                ("rubric_version", "Kontrolversion"),
                ("threshold_note", "Fortolkning af kontrol"),
            ):
                if review.get(key):
                    _paragraph(doc, review[key], label=label)
            for check in review.get("checks", []) or []:
                if not isinstance(check, Mapping):
                    continue
                check_label = check.get("label") or check.get("id") or "Kontrolpunkt"
                state = (
                    "Kræver gennemgang"
                    if check.get("requires_review")
                    else "Ikke markeret til gennemgang"
                )
                if check.get("id") in stale_checks or set(
                    check.get("section_ids", [])
                ).intersection(stale_checks):
                    state = "Tidligere JEV-kontrol — teksten er ændret og kræver nyt kvalitetstjek"
                _paragraph(
                    doc,
                    state,
                    label=_text(check_label),
                    keep_with_next=bool(check.get("section_ids")),
                )
                if check.get("section_ids"):
                    _paragraph(
                        doc,
                        ", ".join(_text(v) for v in check["section_ids"]),
                        label="Berørte afsnit",
                    )
        elif isinstance(review, str):
            _paragraph(doc, review, label="Kontrolresultat")
        sources = ai.get("sources", [])
        if isinstance(sources, list):
            for source in sources:
                if not isinstance(source, Mapping):
                    continue
                identifier = source.get("id") or source.get("source_id") or "Kilde"
                title = (
                    source.get("title")
                    or source.get("filename")
                    or source.get("name")
                    or "Uden titel"
                )
                _heading(doc, f"{_text(identifier)} {_text(title)}", 2)
                source_url = source.get("source_url") or source.get("url")
                # Some historical sources contain only an ID and title. Do
                # not chain such headings through the entire source register.
                doc.paragraphs[-1].paragraph_format.keep_with_next = bool(
                    source_url
                    or any(
                        source.get(key) is not None
                        for key in (
                            "retrieved_at",
                            "document_id",
                            "page",
                            "sha256",
                            "type",
                            "verification_status",
                        )
                    )
                )
                if source_url:
                    _paragraph(doc, source_url, label="Kildeadresse")
                for key, label in (
                    ("retrieved_at", "Hentet"),
                    ("document_id", "Dokument-ID"),
                    ("page", "Side"),
                    ("sha256", "SHA256"),
                    ("type", "Kildetype"),
                    ("verification_status", "Kildestatus"),
                ):
                    if source.get(key) is not None:
                        _paragraph(doc, source[key], label=label)
        limits = ai.get("limitations", [])
        if isinstance(limits, str):
            limits = [limits]
        if isinstance(limits, list) and limits:
            _heading(doc, "Begrænsninger ved AI-udkastet", 2)
            for item in limits:
                _paragraph(doc, item)
    verification = result.legal_verification
    if verification:
        _heading(doc, "Verificering af lovkilder", 2)
        _paragraph(doc, verification.conclusion)
        for receipt in verification.receipts:
            _heading(doc, f"{receipt.citation_id} {receipt.provision}", 3)
            _paragraph(doc, f"{receipt.law} · {receipt.authority}")
            _paragraph(doc, receipt.official_url)
            _paragraph(doc, receipt.message, label="Kontrolresultat")
            _paragraph(doc, receipt.checked_at.isoformat(), label="Kontrolleret")
        for warning in verification.warnings:
            _paragraph(doc, warning, label="Forbehold")
    if not ai and not verification:
        _paragraph(
            doc,
            "Vurderingen er baseret på sagens registrerede oplysninger og den "
            "regelbaserede vurdering. Ingen særskilt kildeverificering er gemt.",
        )


def export_dpia_docx(
    request: DPIAAssessmentRequest, result: DPIAAssessmentResponse
) -> bytes:
    """Return a complete DOCX; reject incomplete or inconsistent snapshots."""
    _validate_snapshot_contract(result)
    guide = getattr(result, "reading_guide", None) or reading_guide(request, result)
    doc = _document(request)
    doc.core_properties.title = _text(f"Konsekvensanalyse for {request.project_name}")[
        :255
    ]
    doc.core_properties.subject = (
        "Databeskyttelsesretlig konsekvensanalyse og risikovurdering"
    )
    doc.core_properties.author = _text(request.organisation)[:255]
    doc.core_properties.identifier = _text(result.id)[:255]
    doc.core_properties.created = result.created_at
    doc.core_properties.modified = result.created_at
    doc.add_paragraph("Konsekvensanalyse vedrørende databeskyttelse", style="Title")
    _paragraph(doc, f"{request.project_name} · {request.organisation}")
    for value, label in (
        (request.department, "Fagområde"),
        (request.processing_version, "Behandlingens version"),
        (request.planned_start_date, "Forventet startdato"),
        (request.planned_start_note, "Bemærkning til opstart"),
        (request.planned_end_date, "Forventet slutdato"),
        (request.planned_end_condition, "Ophørsvilkår"),
    ):
        if value:
            _paragraph(doc, value, label=label)
    _paragraph(
        doc,
        "Databeskyttelsesretlig konsekvensanalyse og risikovurdering efter artikel 35.",
    )
    _paragraph(doc, result.status_label, label="Vurderingsstatus")
    _paragraph(doc, result.id, label="Vurderings-ID")
    _paragraph(doc, getattr(result, "version", 1), label="Version")
    _paragraph(doc, result.created_at.isoformat(), label="Vurderingsdato")
    _paragraph(doc, result.template_version, label="Skabelonversion")
    if result.editorial_revision:
        _paragraph(
            doc,
            "Fagligt redigeret rapportversion. "
            + _editorial_review_note(result)
            + " Se ændringshistorikken under Kilder og vurderingsgrundlag.",
            label="Bearbejdning",
        )
    for field, label in (
        ("case_db_id", "Sags-ID"),
        ("parent_assessment_id", "Forrige vurdering"),
    ):
        if getattr(result, field, None):
            _paragraph(doc, getattr(result, field), label=label)
    _heading(doc, "1 Resumé og scope", new_page=True)
    _heading(doc, "Beslutningsoversigt", 2)
    _paragraph(
        doc,
        guide.get("conclusion") or result.status_label,
        label="Vurderingens aktuelle status",
    )
    _paragraph(
        doc,
        guide.get("approval", {}).get(
            "label", "Godkendelse fremgår ikke af denne rapportversion"
        )
        + ". Hverken et AI-kvalitetstjek eller dokumenterede kontroller udgør en godkendelse.",
        label="Godkendelse",
    )
    for label, values, empty, limit in (
        (
            "Blokerer en godkendelse",
            guide.get("blockers", result.blockers),
            "Ingen blokerende forhold er registreret; faglig stillingtagen kræves stadig.",
            None,
        ),
        (
            "Skal afklares eller dokumenteres",
            guide.get("missing_information", result.missing_information),
            "Ingen manglende oplysninger er registreret i denne version.",
            5,
        ),
        (
            "Næste skridt mod stillingtagen",
            guide.get("next_steps", result.next_steps),
            "Der er ikke registreret næste skridt i denne version.",
            None,
        ),
    ):
        _heading(doc, label, 3)
        _bullet_list(doc, (values[:limit] if limit else values) or [empty])
        if limit and len(values) > limit:
            _paragraph(
                doc,
                f"Oversigten viser de første {limit} af {len(values)} registrerede forhold. Den fulde liste findes i afsnit 7 under Manglende oplysninger.",
            )
    if guide.get("documented_controls"):
        _heading(doc, "Kontroller med registreret evidens – ikke en godkendelse", 3)
        for control in guide["documented_controls"]:
            _structured_text(doc, control["evidence"], label=control["title"])
    _heading(doc, "Resumé", 2)
    _structured_text(doc, result.executive_summary)
    _source_ids(doc, getattr(result, "summary_source_ids", []))
    _paragraph(
        doc, RISK_LEVEL_LABELS[result.risk_level], label="Samlet resterende risiko"
    )
    _heading(doc, "Scope", 2)
    _structured_text(doc, result.scope)
    _recommendations(doc, guide)
    _heading(doc, "Screening af behov for konsekvensanalyse", 2)
    _paragraph(doc, result.screening_conclusion)
    for criterion in result.screening_criteria:
        _paragraph(
            doc,
            criterion.explanation if criterion.matched is None else f"{'Ja' if criterion.matched else 'Nej'}. {criterion.explanation}",
            label=criterion.label,
        )
    sections = {section.id: section for section in result.sections}
    for prefix, title in (
        ("1.", "2 Systematisk beskrivelse af behandlingsaktiviteten"),
        ("2.", "3 Vurdering af lovlighed"),
    ):
        _heading(doc, title, new_page=True)
        for section_id in SECTION_ROWS:
            if not section_id.startswith(prefix):
                continue
            section = sections[section_id]
            _heading(doc, f"{section.id} {section.title}", 2)
            _structured_text(doc, section.text, keep_with_next=True)
            _paragraph(
                doc,
                REVIEW_LABELS.get(section.review_status, section.review_status),
                label="Gennemgang",
                keep_with_next=True,
            )
            _paragraph(
                doc,
                (
                    "Fagligt redigeret. " + _editorial_review_note(result)
                    if f"section:{section.id}"
                    in (result.editorial_revision or {}).get("stale_check_ids", [])
                    else SOURCE_LABELS.get(section.source, section.source)
                ),
                label="Grundlag",
            )
            _source_ids(doc, getattr(section, "source_ids", []))
    _criteria(doc)
    _risk_register(doc, result)
    _heading(doc, "6 Mitigerende foranstaltninger", new_page=True)
    _paragraph(
        doc,
        "De konkrete foranstaltninger er knyttet til hver risiko i afsnit 5. "
        "Her samles sagens oplyste kontroller og dokumentationen for deres gennemførelse.",
    )
    if request.controls:
        _table(
            doc,
            ["Kontrol", "Dokumentationsstatus", "Oplyst evidens"],
            [
                [
                    CONTROL_LABELS[control],
                    (
                        "Verificeret med oplyst evidens"
                        if control in request.verified_controls
                        else "Planlagt eller oplyst uden verificering"
                    ),
                    request.control_evidence.get(control, "Ikke dokumenteret"),
                ]
                for control in request.controls
            ],
            [52, 45, 73],
        )
    else:
        _paragraph(doc, "Der er ikke registreret kontroller i sagens oplysninger.")
    _stakeholders(doc, request, result, guide)
    _provenance(doc, result)
    destination = BytesIO()
    doc.save(destination)
    return destination.getvalue()
